from __future__ import annotations

import copy
import hashlib
import importlib.util
import json
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pypdfium2 as pdfium
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[4]
LATEX = ROOT / "submissions" / "arocmag" / "latex-v2"
OUT = ROOT / "submissions" / "arocmag" / "word-final"
SOURCE = OUT / "source"
FIGURES = OUT / "figures"
FORMULAS = OUT / "formulas"
REPORTS = OUT / "reports"
BASE_DOCX = SOURCE / "official-template-converted.docx"
FINAL_DOCX = OUT / "K-Waay-AROCMAG-submission-stage4.docx"
HELPER_PATH = ROOT / "submissions" / "arocmag" / "word-draft" / "source" / "build_word_manuscript.py"

spec = importlib.util.spec_from_file_location("word_draft_helpers", HELPER_PATH)
h = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(h)


_shared_latex_to_unicode = h.latex_to_unicode


def stage3_latex_to_unicode(src: str, scripts: bool = True) -> str:
    """Preserve the nested j-indices that occur in the first display.

    The shared Stage 1 converter handled digits but flattened ``j_1``,
    ``j_d`` and ``E_j``.  These targeted substitutions retain the intended
    coordinate notation without changing any formula or using raster math.
    """
    if scripts:
        src = src.replace("_{j_1}", "ⱼ₁")
        src = src.replace("_{j_d}", "ⱼ₍d₎")
        src = re.sub(r"_j\b", "ⱼ", src)
    return _shared_latex_to_unicode(src, scripts=scripts)


h.latex_to_unicode = stage3_latex_to_unicode


TITLE_ZH = "K-Waay批次准入中参与方互异条件的形式化分析"
TITLE_EN = "Formal Analysis of the Distinct-Party Condition in K-Waay Batch Admission"

SECTION_ORDER = [
    ("0", "引言", "01-introduction.tex"),
    ("1", "K-Waay批次准入与参与方互异问题", "02-problem.tex"),
    ("2", "批次准入的形式化建模", "03-formal-modeling.tex"),
    ("3", "形式化验证与结果分析", "04-formal-analysis.tex"),
    ("4", "结果讨论与适用范围", "05-discussion.tex"),
    ("5", "结束语", "06-conclusion.tex"),
]

REF_MAP = {
    "tab:identity-coordinate-map": "1",
    "tab:admission-configurations": "2",
    "tab:verification-property-semantics": "3",
    "tab:key-verification-results": "4",
    "fig:two-slot-lifecycle": "1",
    "fig:duplicate-acceptance-trace": "2",
    "eq:abstract-entry": "1",
    "eq:distinct-party-objective": "2",
    "eq:relaxed-duplicate-witness": "3",
    "subsec:configurations-properties": "2.3",
    "subsec:message-results": "3.2",
}

TABLES = {
    "identity-coordinate-map": {
        "caption": "表 1　比较维度与模型映射",
        "widths": [2.2, 2.0, 2.2, 4.2, 5.2],
        "rows": [
            ["比较对象", "论文记号", "Tamarin记号", "作用", "解释边界"],
            ["参与方分量", "$A$", "$A$", "表示模型中的协议主体", "不自动等同于账户或密钥编码"],
            ["发送实例标识", "$\\mathit{oid}$", "$\\mathit{sid}$", "区分规则产生的发送实例", "标识不同不推出参与方不同"],
            ["消息分量", "$m$", "$m$", "表示发送方产生的符号消息", "消息不同不推出参与方不同"],
            ["槽位", "$i$", "线性收集状态", "标识批次内位置", "位置不同不保证条目互异"],
            ["批次与接收方上下文", "$(\\mathit{bid},\\mathit{rst})$", "$(\\mathit{bid},\\mathit{rst})$", "限定共同处理范围", "不表示完整接收方状态"],
        ],
    },
    "admission-configurations": {
        "caption": "表 2　批次准入模型与验证目标",
        "widths": [2.5, 2.2, 2.8, 2.6, 5.7],
        "rows": [
            ["模型", "比较分量", "准入条件", "分析作用", "主要验证对象"],
            ["无互异约束", "无", "无互异要求", "条件移除基线", "来源重复；来源单射性"],
            ["消息互异约束", "消息$m$", "$m_1\\neq m_2$", "替代性检验", "消息互异性；来源单射性；同参与方异消息批次可达性"],
            ["参与方互异约束", "参与方$A$", "$A_1\\neq A_2$", "目标对照", "参与方互异性；有效批次可达性"],
        ],
    },
    "verification-property-semantics": {
        "caption": "表 3　核心验证性质的迹语义",
        "widths": [3.1, 8.0, 4.7],
        "rows": [
            ["性质", "迹语义", "解释边界"],
            ["发送来源对应性", "$\\mathsf{ReceiverAccept}(A,\\mathit{oid},m,\\mathit{bid},\\mathit{rst})@r$推出存在$s<r$使$\\mathsf{Send}(A,\\mathit{oid},m)@s$成立", "断言完整$(A,\\mathit{oid},m)$匹配，不是仅匹配$A$"],
            ["批内来源单射性", "同一$\\mathsf{Send}(A,\\mathit{oid},m)@s$早于两个坐标均为$(A,\\mathit{oid},m,\\mathit{bid},\\mathit{rst})$的接受事件时，推出$r_1=r_2$", "限制一个匹配来源对应的接受事件数，不检查参与方数目"],
            ["批内消息互异性", "同一$(\\mathit{bid},\\mathit{rst})$中若$r_1\\neq r_2$，则该性质断言$m_1\\neq m_2$", "不主张消息跨批次或跨执行全局唯一"],
            ["批内参与方互异性", "同一$(\\mathit{bid},\\mathit{rst})$中若$r_1\\neq r_2$，则该性质要求$A_1\\neq A_2$", "结论只作用于接受事件的参与方投影"],
        ],
    },
    "key-verification-results": {
        "caption": "表 4　关键Tamarin验证结果",
        "widths": [2.2, 6.1, 1.7, 2.1, 3.7],
        "rows": [
            ["模型", "验证性质", "类型", "结果", "论证作用"],
            ["无互异约束", "一个匹配来源对应两个同批接受\none_send_two_accepts_exists", "存在性", "已验证", "给出重复接受的可达执行"],
            ["无互异约束", "批内来源单射性\nreceiver_accept_injective", "全称性", "被反例否定", "表明同一来源可重复对应"],
            ["消息互异约束", "批内消息互异性\naccepted_batch_has_\ndistinct_messages", "全称性", "已验证", "确认模型自身约束"],
            ["消息互异约束", "批内来源单射性\nreceiver_accept_injective", "全称性", "已验证", "排除同一来源重复形状"],
            ["消息互异约束", "同一参与方不同消息批次可达\nsame_party_different_\nmessages_batch_exists", "存在性", "已验证", "支持非替代性结论"],
            ["参与方互异约束", "批内参与方互异性\naccepted_batch_has_\ndistinct_parties", "全称性", "已验证", "保持目标关系"],
            ["参与方互异约束", "有效不同参与方批次可达\ndistinct_party_batch_exists", "存在性", "已验证", "确认有效批次可达"],
        ],
        "note": "注：存在性对应exists-trace，全称性对应all-traces；“已验证”对应verified，“被反例否定”对应falsified。小号英文标识用于追踪相应lemma。",
    },
}


# Final GB/T 7714-2005 renderings.  The venue/publisher/page metadata was
# checked against the cited publisher or conference page and DOI metadata on
# 2026-10-08.  Keeping these renderings here leaves the frozen LaTeX BibTeX
# database untouched while making the Word submission self-contained.
REFERENCE_OVERRIDES = {
    "cremers2023session": {
        "reference": "CREMERS C, JACOMME C, NASKA A. Formal analysis of session-handling in secure messaging: lifting security from sessions to conversations[C]//32nd USENIX Security Symposium (USENIX Security 23). Anaheim, CA: USENIX Association, 2023: 1235-1252.",
        "source": "https://www.usenix.org/conference/usenixsecurity23/presentation/cremers-session-handling",
    },
    "collins2024kwaay": {
        "reference": "COLLINS D, HUGUENIN-DUMITTAN L, NGUYEN N K, et al. K-Waay: fast and deniable post-quantum X3DH without ring signatures[C]//33rd USENIX Security Symposium (USENIX Security 24). Philadelphia, PA: USENIX Association, 2024: 433-450.",
        "source": "https://www.usenix.org/conference/usenixsecurity24/presentation/collins",
    },
    "collins2024kwaayfull": {
        "reference": "COLLINS D, HUGUENIN-DUMITTAN L, NGUYEN N K, et al. K-Waay: fast and deniable post-quantum X3DH without ring signatures[EB/OL]. Cryptology ePrint Archive, Paper 2024/120, 2024[2026-10-08]. https://eprint.iacr.org/2024/120.",
        "source": "https://eprint.iacr.org/2024/120",
    },
    "cohngordon2020signal": {
        "reference": "COHN-GORDON K, CREMERS C, DOWLING B, et al. A formal security analysis of the Signal messaging protocol[J]. Journal of Cryptology, 2020, 33(4): 1914-1983. DOI:10.1007/s00145-020-09360-1.",
        "source": "https://link.springer.com/article/10.1007/s00145-020-09360-1",
    },
    "bhargavan2024pqxdh": {
        "reference": "BHARGAVAN K, JACOMME C, KIEFER F, et al. Formal verification of the PQXDH post-quantum key agreement protocol for end-to-end secure messaging[C]//33rd USENIX Security Symposium (USENIX Security 24). Philadelphia, PA: USENIX Association, 2024: 469-486.",
        "source": "https://www.usenix.org/conference/usenixsecurity24/presentation/bhargavan",
    },
    "meier2013tamarin": {
        "reference": "MEIER S, SCHMIDT B, CREMERS C, et al. The TAMARIN prover for the symbolic analysis of security protocols[C]//Computer Aided Verification: CAV 2013. Berlin, Heidelberg: Springer, 2013: 696-701. DOI:10.1007/978-3-642-39799-8_48.",
        "source": "https://link.springer.com/chapter/10.1007/978-3-642-39799-8_48",
    },
    "lupetti2006names": {
        "reference": "LUPETTI S, DILLEMA F W, STABELL-KULO T. Names in cryptographic protocols[C]//Proceedings of the 4th International Workshop on Security in Information Systems. Paphos, Cyprus: SciTePress, 2006: 185-194. DOI:10.5220/0002484701850194.",
        "source": "https://www.scitepress.org/PublishedPapers/2006/24847/",
    },
    "lowe1997hierarchy": {
        "reference": "LOWE G. A hierarchy of authentication specifications[C]//Proceedings of the 10th Computer Security Foundations Workshop. Rockport, MA: IEEE Computer Society Press, 1997: 31-43. DOI:10.1109/CSFW.1997.596782.",
        "source": "https://doi.org/10.1109/CSFW.1997.596782",
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def extract_abstracts() -> tuple[str, str, str, str]:
    text = (LATEX / "sections" / "00-abstract.tex").read_text(encoding="utf-8")
    zh = re.search(r"\\begin\{abstract\}\s*(.+?)\s*\\noindent\\textbf\{关键词：\}", text, re.S).group(1)
    zh_kw = re.search(r"\\textbf\{关键词：\}(.+?)\\end\{abstract\}", text, re.S).group(1)
    en = re.search(r"\\noindent\s*K-Waay(.+?)\\noindent\\textbf\{Key words:\}", text, re.S).group(0)
    en = re.sub(r"^\\noindent\s*", "", en)
    en = re.sub(r"\s*\\noindent\\textbf\{Key words:\}\s*$", "", en)
    en_kw = re.search(r"\\textbf\{Key words:\}\s*(.+?)\s*$", text, re.S).group(1)
    normalize = lambda value: re.sub(r"\s+", " ", value).strip()
    def manuscript_text(value: str) -> str:
        return (
            normalize(value)
            .replace("\\batchreceive", "BatchReceive")
            .replace("\\receiveraccept", "ReceiverAccept")
            .replace("\\sendaction", "Send")
            .replace("\\kwaay", "K-Waay")
            .replace("---", "—")
            .replace("--", "–")
        )

    return manuscript_text(zh), normalize(zh_kw), manuscript_text(en), normalize(en_kw)


def citation_order() -> list[str]:
    order: list[str] = []
    for _, _, name in SECTION_ORDER:
        text = (LATEX / "sections" / name).read_text(encoding="utf-8")
        for match in re.finditer(r"\\cite(?:\[[^\]]+\])?\{([^}]+)\}", text):
            for key in match.group(1).split(","):
                key = key.strip()
                if key and key not in order:
                    order.append(key)
    return order


def replace_citations(text: str, cite_map: dict[str, int]) -> str:
    def repl(match: re.Match) -> str:
        nums = [cite_map[key.strip()] for key in match.group(1).split(",")]
        return "[" + ",".join(str(value) for value in nums) + "]"

    return re.sub(r"\\cite(?:\[[^\]]+\])?\{([^}]+)\}", repl, text)


def tex_text(text: str, cite_map: dict[str, int]) -> str:
    text = replace_citations(text, cite_map)
    text = re.sub(r"\\eqref\{([^}]+)\}", lambda m: f"({REF_MAP[m.group(1)]})", text)
    text = re.sub(r"\\ref\{([^}]+)\}", lambda m: REF_MAP[m.group(1)], text)
    text = text.replace("\\kwaay", "K-Waay").replace("\\batchreceive", "BatchReceive")
    text = text.replace("\\receiveraccept", "ReceiverAccept").replace("\\sendaction", "Send")
    text = text.replace("\\distinctparty", "\\mathsf{DistinctPartyPerBatch}")
    text = text.replace("\\oid", "\\mathit{oid}")
    # The shared OMML helper understands \not\Rightarrow but not the AMS
    # alias \nRightarrow used by the current manuscript.
    text = text.replace("\\nRightarrow", "\\not\\Rightarrow")
    text = text.replace("\\textasciitilde", "__LITERAL_TILDE__").replace("\\%", "%").replace("\\&", "&")
    text = text.replace("\\newline", "\n").replace("---", "—").replace("--", "–")
    text = re.sub(r"\\(?:path|texttt|textsf|textsc|emph|textbf)\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\(?:path|texttt|textsf|textsc|emph|textbf)\{([^{}]*)\}", r"\1", text)
    text = text.replace("~", " ").replace("__LITERAL_TILDE__", "~").replace("\\xspace", "")
    return re.sub(r"[ \t]+", " ", text).strip()


def add_rich_paragraph(doc: Document, text: str, cite_map: dict[str, int], indent: bool = True):
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "71"))
    # The official body style may resolve to distributed alignment in WPS.
    # Paragraphs containing long lemma identifiers use left alignment so the
    # surrounding Chinese is not stretched; ordinary prose remains justified.
    h.set_paragraph_metrics(paragraph, first_indent=indent)
    converted = tex_text(text, cite_map)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
        if re.search(r"[a-z][a-z0-9]*(?:_[a-z0-9]+){2,}", converted)
        else WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    h.add_inline(paragraph, converted, cite_map)
    return paragraph


def add_list_paragraph(doc: Document, label: str, text: str, cite_map: dict[str, int]):
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "71"))
    paragraph.paragraph_format.left_indent = Pt(22)
    paragraph.paragraph_format.first_line_indent = Pt(-22)
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(12)
    run = paragraph.add_run(label)
    h.set_run_font(run, size=9, bold=True)
    converted = tex_text(text, cite_map)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
        if re.search(r"[a-z][a-z0-9]*(?:_[a-z0-9]+){2,}", converted)
        else WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    h.add_inline(paragraph, converted, cite_map)
    return paragraph


def clean_formula(raw: str) -> str:
    raw = re.sub(r"\\label\{[^}]+\}", "", raw)
    raw = raw.replace("\\oid", "\\mathit{oid}").replace("\\distinctparty", "\\mathsf{DistinctPartyPerBatch}")
    return raw.strip()


def add_formula(doc: Document, raw: str, number: int | None):
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "74"))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(8.1), WD_TAB_ALIGNMENT.RIGHT)
    lines = [item.strip() for item in re.split(r"\\\\", clean_formula(raw)) if item.strip()]
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        paragraph._p.append(h.omath(line))
    if number is not None:
        run = paragraph.add_run(f"\t({number})")
        h.set_run_font(run, size=9)
    return paragraph


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell(cell, value: str, cite_map: dict[str, int], bold: bool, size: float) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(9.5)
    for index, line in enumerate(value.split("\n")):
        if index:
            paragraph.add_run().add_break()
        before = len(paragraph.runs)
        h.add_inline(paragraph, tex_text(line, cite_map), cite_map, size=size)
        for run in paragraph.runs[before:]:
            h.set_run_font(run, size=size, bold=bold if index == 0 else False)
            if index > 0:
                run.font.size = Pt(max(size - 0.5, 6.5))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table_block(doc: Document, name: str, cite_map: dict[str, int]) -> None:
    data = TABLES[name]
    one_col = doc.add_section(WD_SECTION.CONTINUOUS)
    h.set_section_geometry(one_col)
    h.set_columns(one_col, 1)
    caption = doc.add_paragraph(style=h.style_by_id(doc, "77"))
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    h.add_text_run(caption, data["caption"], bold=True, size=8)

    rows = data["rows"]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col_index, width in enumerate(data["widths"]):
        table.columns[col_index].width = Cm(width)
    none = {"val": "nil"}
    rule = {"val": "single", "sz": "6", "space": "0", "color": "000000"}
    for row_index, row in enumerate(rows):
        prevent_row_split(table.rows[row_index])
        if row_index == 0:
            set_repeat_table_header(table.rows[row_index])
        table.rows[row_index].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.width = Cm(data["widths"][col_index])
            set_cell(cell, value, cite_map, row_index == 0, 7.0 if len(row) >= 5 else 7.5)
            h.set_cell_border(
                cell,
                top=rule if row_index == 0 else none,
                bottom=rule if row_index in (0, len(rows) - 1) else none,
                left=none,
                right=none,
                insideH=none,
                insideV=none,
            )
    if data.get("note"):
        note = doc.add_paragraph(style=h.style_by_id(doc, "80"))
        note.paragraph_format.space_before = Pt(1)
        note.paragraph_format.space_after = Pt(2)
        h.add_text_run(note, data["note"], size=7.5)
    two_col = doc.add_section(WD_SECTION.CONTINUOUS)
    h.set_section_geometry(two_col)
    h.set_columns(two_col, 2)


def render_figure_pdf(number: int) -> Path:
    source = FIGURES / f"fig{number}-preview.pdf"
    target = FIGURES / f"fig{number}-word-stage4.png"
    pdf = pdfium.PdfDocument(str(source))
    page = pdf[0]
    bitmap = page.render(scale=5)
    bitmap.to_pil().convert("RGB").save(target, dpi=(360, 360))
    bitmap.close()
    page.close()
    pdf.close()
    with Image.open(target) as image:
        # The standalone preview carries its own caption (and Figure 2 note).
        # The Word manuscript uses editable caption/note text, so retain only
        # the frozen TikZ drawing area from the preview artifact.
        bottom_ratio = 0.91 if number == 1 else 0.775
        cropped = image.crop((20, 20, image.width - 20, int(image.height * bottom_ratio)))
        cropped.save(target, dpi=(360, 360))
    return target


def add_figure(doc: Document, number: int) -> None:
    titles = {
        1: "图 1　两槽批次准入的共同生命周期",
        2: "图 2　无互异约束模型中的重复接受可达执行",
    }
    image_path = FIGURES / f"fig{number}-word-stage4.png"
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "77"))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(image_path), width=Cm(7.9))
    paragraph.add_run().add_break()
    h.add_text_run(paragraph, titles[number], size=8)
    if number == 2:
        note = doc.add_paragraph(style=h.style_by_id(doc, "80"))
        note.paragraph_format.space_before = Pt(1)
        note.paragraph_format.space_after = Pt(2)
        h.add_text_run(
            note,
            "依据Tamarin模型及独立复核的可达执行重构，仅显示与所示完整元组相关的关键事件；该元组的匹配Send唯一，执行中可能存在与其不匹配的其他Send。",
            size=7.5,
        )


def add_front_matter(doc: Document) -> None:
    zh_abstract, zh_keywords, en_abstract, en_keywords = extract_abstracts()
    title = doc.add_paragraph(style=h.style_by_id(doc, "40"))
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    h.add_text_run(title, TITLE_ZH, bold=True, size=16)

    for style_id, text, size in [
        ("41", "作者信息待补", 12),
        ("43", "（作者单位、城市、邮编待补）", 9),
    ]:
        paragraph = doc.add_paragraph(style=h.style_by_id(doc, style_id))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        h.add_text_run(paragraph, text, size=size)

    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "45"))
    h.set_paragraph_metrics(paragraph, first_indent=False)
    h.add_text_run(paragraph, "摘  要：", bold=True)
    h.add_text_run(paragraph, zh_abstract)
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "45"))
    h.set_paragraph_metrics(paragraph, first_indent=False)
    h.add_text_run(paragraph, "关键词：", bold=True)
    h.add_text_run(paragraph, zh_keywords)
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "45"))
    h.set_paragraph_metrics(paragraph, first_indent=False)
    h.add_text_run(paragraph, "中图分类号：待确认　　文献标志码：待确认　　文章编号：待定", bold=True)

    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "42"))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    h.add_text_run(paragraph, TITLE_EN, bold=True, size=12)
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "56"))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.add_text_run(paragraph, "Author information pending", size=10.5)
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "43"))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.add_text_run(paragraph, "(Affiliation, city, postal code, and country pending)", size=9)
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "45"))
    h.set_paragraph_metrics(paragraph, first_indent=False, line=11)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h.add_text_run(paragraph, "Abstract: ", bold=True, size=8.5)
    h.add_text_run(paragraph, en_abstract, size=8.5)
    paragraph = doc.add_paragraph(style=h.style_by_id(doc, "45"))
    h.set_paragraph_metrics(paragraph, first_indent=False, line=11)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h.add_text_run(paragraph, "Key words: ", bold=True, size=8.5)
    h.add_text_run(paragraph, en_keywords, size=8.5)


def consume_environment(lines: list[str], start: int, end_marker: str) -> tuple[str, int]:
    content: list[str] = []
    index = start
    while index < len(lines) and end_marker not in lines[index]:
        content.append(lines[index])
        index += 1
    return "\n".join(content), index + 1


def add_section_content(doc: Document, path: Path, cite_map: dict[str, int], formula_records: list[dict]) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    numbered_formula_count = len([item for item in formula_records if item["number"] is not None])
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith("%") or line.startswith("\\label{"):
            index += 1
            continue
        match = re.match(r"\\subsection\{(.+)\}", line)
        if match:
            h.add_heading(doc, match.group(1), 2)
            index += 1
            continue
        if line.startswith("\\begin{figure}"):
            block, index = consume_environment(lines, index + 1, "\\end{figure}")
            number = 1 if "fig1-two-slot-lifecycle" in block else 2
            add_figure(doc, number)
            continue
        table_match = re.match(r"\\input\{tables/([^}]+)\}", line)
        if table_match:
            add_table_block(doc, table_match.group(1), cite_map)
            index += 1
            continue
        if line.startswith("\\begin{equation}"):
            raw, index = consume_environment(lines, index + 1, "\\end{equation}")
            numbered_formula_count += 1
            add_formula(doc, raw, numbered_formula_count)
            formula_records.append({"source": path.name, "number": numbered_formula_count, "latex": clean_formula(raw)})
            continue
        if line == "\\[":
            raw, index = consume_environment(lines, index + 1, "\\]")
            add_formula(doc, raw, None)
            formula_records.append({"source": path.name, "number": None, "latex": clean_formula(raw)})
            continue
        if line.startswith("\\begin{enumerate}"):
            rq_mode = "RQ" in line
            items: list[str] = []
            current: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("\\end{enumerate}"):
                candidate = lines[index].strip()
                if candidate.startswith("\\item"):
                    if current:
                        items.append(" ".join(current))
                    current = [re.sub(r"^\\item\s*", "", candidate)]
                elif candidate:
                    current.append(candidate)
                index += 1
            if current:
                items.append(" ".join(current))
            for item_index, item in enumerate(items, 1):
                add_list_paragraph(doc, f"RQ{item_index}：" if rq_mode else f"{item_index}) ", item, cite_map)
            index += 1
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                index += 1
                break
            if (
                candidate.startswith("%")
                or candidate.startswith("\\subsection{")
                or candidate.startswith("\\label{")
                or candidate.startswith("\\begin{figure}")
                or candidate.startswith("\\input{tables/")
                or candidate.startswith("\\begin{equation}")
                or candidate.startswith("\\begin{enumerate}")
                or candidate == "\\["
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        add_rich_paragraph(doc, " ".join(paragraph_lines), cite_map)


def replace_story(destination, source) -> None:
    element = destination._element
    for child in list(element):
        element.remove(child)
    for child in source._element:
        element.append(copy.deepcopy(child))


def replace_paragraph_text(paragraph, value: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(value)
        return
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


def apply_header_footer(doc: Document, template: Document) -> None:
    first = doc.sections[0]
    source = template.sections[0]
    first.different_first_page_header_footer = True
    first.header.is_linked_to_previous = False
    first.first_page_header.is_linked_to_previous = False
    first.footer.is_linked_to_previous = False
    first.first_page_footer.is_linked_to_previous = False
    replace_story(first.header, source.header)
    replace_story(first.first_page_header, source.first_page_header)
    replace_story(first.footer, source.footer)
    replace_story(first.first_page_footer, source.first_page_footer)
    if first.first_page_header.paragraphs:
        replace_paragraph_text(
            first.first_page_header.paragraphs[0],
            "\n投稿\t计算机应用研究\t修改日期：2026/10/09",
        )

    footer = first.first_page_footer
    while len(footer.paragraphs) < 4:
        footer.add_paragraph()
    replace_paragraph_text(footer.paragraphs[0], "——————————")
    replace_paragraph_text(footer.paragraphs[1], "　　基金项目：待补")
    replace_paragraph_text(footer.paragraphs[2], "　　作者简介：待补")
    replace_paragraph_text(footer.paragraphs[3], "　　通信作者及电子邮箱：待补")

    for section in doc.sections[1:]:
        section.different_first_page_header_footer = False
        section.header.is_linked_to_previous = True
        section.first_page_header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True
        section.first_page_footer.is_linked_to_previous = True


def add_references(doc: Document, order: list[str], bib: dict) -> list[dict]:
    heading = doc.add_paragraph(style=h.style_by_id(doc, "82"))
    heading.paragraph_format.keep_with_next = True
    h.add_text_run(heading, "参考文献", bold=True, size=10.5)
    records: list[dict] = []
    for number, key in enumerate(order, 1):
        # WPS does not balance the final two-column section automatically.
        # Keep reference 1 after the conclusion and begin reference 2 in
        # the right column, without changing content or reference order.
        if number == 2:
            break_paragraph = doc.add_paragraph(style=h.style_by_id(doc, "71"))
            break_paragraph.paragraph_format.space_before = Pt(0)
            break_paragraph.paragraph_format.space_after = Pt(0)
            break_paragraph.add_run().add_break(WD_BREAK.COLUMN)
        verified = REFERENCE_OVERRIDES[key]
        formatted = f"[{number}] {verified['reference']}"
        paragraph = doc.add_paragraph(style=h.style_by_id(doc, "84"))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.first_line_indent = Pt(-18)
        paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(9.5)
        h.add_text_run(paragraph, re.sub(r"^\[\d+\]\s*", "", formatted), size=8)
        records.append({
            "number": number,
            "key": key,
            "reference": formatted,
            "verified_source": verified["source"],
            "source_metadata_pending": False,
        })
    return records


def normalize_styles(doc: Document) -> None:
    for style_id, east_asia, latin, size in [
        ("40", "黑体", "Arial", 16),
        ("41", "宋体", "Times New Roman", 12),
        ("43", "楷体", "Times New Roman", 9),
        ("45", "楷体", "Times New Roman", 9),
        ("42", "宋体", "Times New Roman", 12),
        ("56", "宋体", "Times New Roman", 10.5),
        ("65", "黑体", "Arial", 10.5),
        ("66", "黑体", "Arial", 9),
        ("71", "宋体", "Times New Roman", 9),
        ("74", "宋体", "Cambria Math", 9),
        ("77", "宋体", "Times New Roman", 8),
        ("80", "宋体", "Times New Roman", 7.5),
        ("82", "黑体", "Arial", 10.5),
        ("84", "楷体", "Times New Roman", 8),
    ]:
        style = h.style_by_id(doc, style_id)
        style.font.name = latin
        style.font.size = Pt(size)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def write_source_records(order: list[str], references: list[dict], formulas: list[dict]) -> None:
    (FORMULAS / "formulas-stage4.json").write_text(json.dumps(formulas, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (SOURCE / "citation-order-stage4.json").write_text(json.dumps(references, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    git_head = subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True, encoding="utf-8"
    ).strip()
    scientific_sources = {
        str(path.relative_to(ROOT)).replace("\\", "/"): sha256(path)
        for path in sorted((LATEX / "sections").glob("*.tex"))
    }
    manifest = {
        "scientific_baseline_head": git_head,
        "scientific_sources": scientific_sources,
        "official_template_sha256": "BAF9FCE8B2BC56F2320C8BF4541546E43F48A8F3A1C773B40667D7B2A7B872CB",
        "working_docx": str(FINAL_DOCX.relative_to(ROOT)).replace("\\", "/"),
        "working_docx_sha256": sha256(FINAL_DOCX),
        "citation_order": order,
        "numbered_formulas": len([item for item in formulas if item["number"] is not None]),
        "unnumbered_formulas": len([item for item in formulas if item["number"] is None]),
        "tables": 4,
        "figures": 2,
        "figure_format": "PNG, 360 dpi, generated from frozen PDF previews",
        "equation_format": "editable OMML",
        "mathtype_status": "MATHTYPE_CONVERSION_BLOCKED",
        "vector_graphics_status": "VECTOR_GRAPHICS_PENDING",
        "reference_standard": "GB/T 7714-2005",
        "official_requirement_check_date": "2026-10-09",
        "official_requirement_pages": [
            "https://www.arocmag.cn/info/instruction/template",
            "https://www.arocmag.cn/info/instruction/abstract-instruction",
            "https://www.arocmag.cn/info/instruction/instructions",
        ],
    }
    (SOURCE / "build-manifest-stage4.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def remove_orphaned_template_ole_parts(path: Path) -> None:
    """Drop unused OLE objects retained by the converted official template.

    The document body has no ``o:OLEObject`` element after reconstruction, but
    python-docx preserves two now-unreferenced template OLE binaries and their
    document relationships.  Removing those stale package parts avoids
    misreporting them as MathType equations.
    """
    temp = path.with_suffix(".cleaning.docx")
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    ET.register_namespace("", rel_ns)
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            name = item.filename
            if name.startswith("word/embeddings/") and name.lower().endswith(".bin"):
                continue
            data = source.read(name)
            if name == "word/_rels/document.xml.rels":
                root = ET.fromstring(data)
                for relationship in list(root):
                    if relationship.get("Type", "").endswith("/oleObject"):
                        root.remove(relationship)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif name == "[Content_Types].xml":
                root = ET.fromstring(data)
                for node in list(root):
                    if node.tag == f"{{{ct_ns}}}Default" and node.get("Extension", "").lower() == "bin":
                        root.remove(node)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
    temp.replace(path)


def build() -> None:
    if sha256(ROOT / "submissions" / "arocmag" / "official" / "templates" / "投稿编辑模板-计算机应用研究.doc") != "BAF9FCE8B2BC56F2320C8BF4541546E43F48A8F3A1C773B40667D7B2A7B872CB":
        raise RuntimeError("Official template changed; redistillation is required")
    for folder in (SOURCE, FIGURES, FORMULAS, REPORTS):
        folder.mkdir(parents=True, exist_ok=True)
    render_figure_pdf(1)
    render_figure_pdf(2)

    order = citation_order()
    bib = h.parse_bib((LATEX / "bibliography" / "references.bib").read_text(encoding="utf-8"))
    missing = [key for key in order if key not in bib]
    if missing:
        raise RuntimeError("Missing bibliography entries: " + ", ".join(missing))
    cite_map = {key: number for number, key in enumerate(order, 1)}

    template = Document(BASE_DOCX)
    doc = Document(BASE_DOCX)
    h.clear_document_body(doc)
    normalize_styles(doc)
    first = doc.sections[0]
    h.set_section_geometry(first)
    h.set_columns(first, 1)
    add_front_matter(doc)

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    h.set_section_geometry(body)
    h.set_columns(body, 2)
    formula_records: list[dict] = []
    for _, heading, filename in SECTION_ORDER:
        h.add_heading(doc, heading, 1)
        add_section_content(doc, LATEX / "sections" / filename, cite_map, formula_records)
    references = add_references(doc, order, bib)
    apply_header_footer(doc, template)

    properties = doc.core_properties
    properties.title = TITLE_EN
    properties.subject = "Anonymous AROCMAG Stage 4 submission-compliance manuscript"
    properties.author = "Anonymous Author(s)"
    properties.last_modified_by = "Anonymous Author(s)"
    properties.keywords = "K-Waay; batch admission; distinct-party condition; Tamarin; formal verification"
    properties.comments = "Official-template-derived anonymous Stage 4 manuscript; submission-compliance build synchronized from the current latex-v2 source."
    doc.save(FINAL_DOCX)
    remove_orphaned_template_ole_parts(FINAL_DOCX)
    write_source_records(order, references, formula_records)
    print(json.dumps({
        "docx": str(FINAL_DOCX),
        "sha256": sha256(FINAL_DOCX),
        "citations": order,
        "references": len(references),
        "numbered_formulas": len([item for item in formula_records if item["number"] is not None]),
        "unnumbered_formulas": len([item for item in formula_records if item["number"] is None]),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
