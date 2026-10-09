from __future__ import annotations

import json
import importlib.util
import re
import sys
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[4]
WORD_ROOT = ROOT / "submissions" / "arocmag" / "word-final"
DOCX = WORD_ROOT / "K-Waay-AROCMAG-submission-stage4.docx"
OUT = WORD_ROOT / "qa" / "stage4-content-audit.json"
BUILDER_PATH = WORD_ROOT / "source" / "build_word_stage4.py"

spec = importlib.util.spec_from_file_location("word_stage3_builder", BUILDER_PATH)
b = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(b)


def norm(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u200b", "")).strip()


def package_visible_text(path: Path) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name == "word/document.xml" or re.fullmatch(r"word/(header|footer)\d+\.xml", name):
                root = etree.fromstring(archive.read(name))
                chunks.extend(root.xpath("//*[local-name()='t']/text()"))
    return "\n".join(chunks)


def main() -> None:
    doc = Document(DOCX)
    visible = package_visible_text(DOCX).replace("\u200b", "")
    compact = norm(visible)
    squash = re.sub(r"\s+", "", visible)
    paragraph_texts = [norm(paragraph.text) for paragraph in doc.paragraphs]

    expected_zh, _, expected_en, _ = b.extract_abstracts()
    zh_para = next(text for text in paragraph_texts if text.startswith("摘 要：") or text.startswith("摘 要：") or text.startswith("摘  要："))
    en_para = next(text for text in paragraph_texts if text.startswith("Abstract:"))
    zh_actual = re.sub(r"^摘\s*要：", "", zh_para).strip()
    en_actual = en_para.removeprefix("Abstract:").strip()

    forbidden_patterns = {
        "latex_cite": r"\\cite(?:\[|\{)",
        "latex_ref": r"\\(?:eqref|ref)\{",
        "latex_environment": r"\\(?:begin|end|input)\{",
        "latex_project_macros": r"\\(?:batchreceive|receiveraccept|sendaction|oid|kwaay)\b",
        "todo": r"\bTODO\b",
        "local_path": r"(?:[A-Za-z]:\\|kwaay-formal|word-draft)",
        "github_internal": r"github\.com/.+/(?:blob|tree)/",
        "misleading_source_wording": r"合法来源|合法匹配来源|legitimate matching origin|legitimate origin",
    }
    forbidden_hits = {name: re.findall(pattern, visible, flags=re.I) for name, pattern in forbidden_patterns.items()}

    anchors = [
        "same_party_different_messages_batch_exists",
        "全部14项结果仍为13项验证通过、1项被反例否定",
        "Pτ⇒Mτ",
        "Mτ⇔Iτ",
        "Mτ⇏Pτ",
        "而非新增的Tamarin lemma或证明器结果",
        "机器证据固定为两个槽位",
        "不能直接解释为具体协议攻击",
        "结论限于两槽准入抽象和ReceiverAccept事件",
        "该元组的匹配Send唯一，执行中可能存在与其不匹配的其他Send",
        "模型内匹配发送来源",
    ]
    anchor_status = {anchor: re.sub(r"\s+", "", anchor) in squash for anchor in anchors}

    table_shapes = [[len(table.rows), len(table.columns)] for table in doc.tables]
    heading_counts = {
        "level1": sum(paragraph.style.style_id == "65" for paragraph in doc.paragraphs),
        "level2": sum(paragraph.style.style_id == "66" for paragraph in doc.paragraphs),
    }
    omath_count = len(doc.element.xpath(".//*[local-name()='oMath']"))
    inline_shapes = len(doc.inline_shapes)
    numbered_equations = [number for number in (1, 2, 3) if any(f"({number})" in text for text in paragraph_texts)]
    cited = set()
    for group in re.findall(r"\[([0-9,，\s]+)\]", compact):
        cited.update(int(value) for value in re.findall(r"\d+", group))
    citation_numbers = sorted(number for number in cited if 1 <= number <= 8)
    formula_records = json.loads((WORD_ROOT / "formulas" / "formulas-stage4.json").read_text(encoding="utf-8"))
    formula_paragraphs = [paragraph for paragraph in doc.paragraphs if paragraph.style.style_id == "74"]
    actual_display_formulas = [
        "\n".join(paragraph._p.xpath('.//*[local-name()="oMath"]//*[local-name()="t"]/text()'))
        for paragraph in formula_paragraphs
    ]
    expected_display_formulas = []
    for record in formula_records:
        lines = [item.strip() for item in re.split(r"\\\\", b.clean_formula(record["latex"])) if item.strip()]
        expected_display_formulas.append("\n".join(b.h.latex_to_unicode(line) for line in lines))
    formula_contents_exact = [
        re.sub(r"\s+", "", actual) == re.sub(r"\s+", "", expected)
        for actual, expected in zip(actual_display_formulas, expected_display_formulas)
    ]
    reference_texts = [
        text for paragraph, text in zip(doc.paragraphs, paragraph_texts)
        if paragraph.style.style_id == "84" and text
    ]
    expected_references = [b.REFERENCE_OVERRIDES[key]["reference"] for key in b.citation_order()]

    relation_paragraph = next((text for text in paragraph_texts if "机器结果直接支持" in text), "")
    result_paragraph = next((text for text in paragraph_texts if "全部14项结果仍为" in text), "")
    modeling_scope_paragraph = next((text for text in paragraph_texts if "模型以持久事实" in text), "")
    semantic_context_checks = {
        "machine_witness_vs_manual_derivation": all(token in relation_paragraph for token in [
            "same_party_different_messages_batch_exists", "机器结果直接支持", "人工迹语义推导", "而非新增的Tamarin lemma"
        ]),
        "fourteen_result_conservation": all(token in result_paragraph for token in ["14项", "13项验证通过", "1项被反例否定"]),
        "receiver_accept_scope": all(token in modeling_scope_paragraph for token in [
            "!Sent(A,oid,m)", "没有实现签名验证", "因而不能推出密钥泄露"
        ]),
    }

    core = doc.core_properties
    metadata = {
        "title": core.title,
        "subject": core.subject,
        "author": core.author,
        "last_modified_by": core.last_modified_by,
    }
    checks = {
        "title_exact": TITLE_ZH == paragraph_texts[0],
        "chinese_abstract_exact": norm(expected_zh) == norm(zh_actual),
        "english_abstract_exact": norm(expected_en) == norm(en_actual),
        "keywords_present": "关键词：K-Waay；批次准入；批内参与方互异条件；Tamarin；形式化验证" in squash,
        "english_keywords_present": "Key words: K-Waay; batch admission; distinct-party condition; Tamarin; formal verification" in compact,
        "all_scientific_anchors": all(anchor_status.values()),
        "four_tables_exact_shapes": table_shapes == [[6, 5], [4, 5], [5, 3], [8, 5]],
        "two_figures": inline_shapes == 2,
        "three_numbered_equations": numbered_equations == [1, 2, 3],
        "editable_omml_present": omath_count > 0,
        "three_numbered_and_three_unnumbered_formulas": (
            len([item for item in formula_records if item["number"] is not None]) == 3
            and len([item for item in formula_records if item["number"] is None]) == 3
        ),
        "six_display_formula_contents_exact": len(formula_contents_exact) == 6 and all(formula_contents_exact),
        "all_eight_citations_present": citation_numbers == list(range(1, 9)),
        "all_eight_references_exact": reference_texts == expected_references,
        "heading_counts": heading_counts == {"level1": 6, "level2": 13},
        "semantic_context_checks": all(semantic_context_checks.values()),
        "no_forbidden_visible_text": all(not values for values in forbidden_hits.values()),
        "no_unconverted_latex_em_dashes": "---" not in visible,
        "anonymous_metadata": metadata["author"] == "Anonymous Author(s)" and metadata["last_modified_by"] == "Anonymous Author(s)",
        "stage4_metadata": metadata["subject"] == "Anonymous AROCMAG Stage 4 submission-compliance manuscript",
    }
    result = {
        "checks": checks,
        "all_pass": all(checks.values()),
        "anchor_status": anchor_status,
        "forbidden_hits": forbidden_hits,
        "table_shapes": table_shapes,
        "heading_counts": heading_counts,
        "omath_count": omath_count,
        "inline_shapes": inline_shapes,
        "numbered_equations": numbered_equations,
        "citation_numbers": citation_numbers,
        "semantic_context_checks": semantic_context_checks,
        "formula_records": formula_records,
        "display_formula_contents_exact": formula_contents_exact,
        "references_exact": reference_texts == expected_references,
        "metadata": metadata,
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["all_pass"]:
        raise SystemExit(1)


TITLE_ZH = "K-Waay批次准入中参与方互异条件的形式化分析"


if __name__ == "__main__":
    main()
