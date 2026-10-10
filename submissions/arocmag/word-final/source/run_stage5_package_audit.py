from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree
from PIL import Image
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[4]
WORD_ROOT = ROOT / "submissions" / "arocmag" / "word-final"
PACKAGE = WORD_ROOT / "stage5-submission-package"
STAGE4_DOCX = WORD_ROOT / "K-Waay-AROCMAG-submission-stage4.docx"
STAGE5_DOCX = PACKAGE / "K-Waay-AROCMAG-submission-stage5-author-confirmation.docx"
STAGE4_PDF = WORD_ROOT / "qa" / "stage4-render-final" / "K-Waay-AROCMAG-submission-stage4-verified.pdf"
STAGE5_PDF = PACKAGE / "K-Waay-AROCMAG-submission-stage5-preview.pdf"
FORMULA_CHECKLIST = PACKAGE / "FORMULA_VERIFICATION_CHECKLIST.csv"
RENDER_DIR = WORD_ROOT / "qa" / "stage5-render-final"
AUDIT_OUT = WORD_ROOT / "qa" / "stage5-package-audit.json"
MANIFEST_OUT = PACKAGE / "PACKAGE_MANIFEST.json"

EXPECTED_PACKAGE_FILES = [
    "AUTHOR_INFORMATION_COLLECTION_FORM.md",
    "FORMULA_VERIFICATION_CHECKLIST.csv",
    "FORMULA_VERIFICATION_GUIDE.md",
    "K-Waay-AROCMAG-submission-stage5-author-confirmation.docx",
    "K-Waay-AROCMAG-submission-stage5-preview.pdf",
    "PACKAGE_README.md",
    "SUBMISSION_COMPLIANCE_REPORT.md",
    "SUBMISSION_SYSTEM_OPERATION_CHECKLIST.md",
]
EXPECTED_FIGURE_HASHES = [
    "AE6C5800ECFA64AB86200A6C73E25AA9248BA584825B9EEA8D0BC804E25C12AB",
    "8A4804F82DC382DE777B726D21827703DA84C7E7920311F703619EB48DA06992",
]

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "o": "urn:schemas-microsoft-com:office:office",
}


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def normalized_pdf_text(path: Path) -> list[str]:
    return [re.sub(r"\s+", "", page.extract_text() or "") for page in PdfReader(str(path)).pages]


def package_manifest() -> dict:
    files = []
    for path in sorted(PACKAGE.iterdir(), key=lambda item: item.name.lower()):
        if path.is_file() and path.name != MANIFEST_OUT.name:
            files.append({
                "name": path.name,
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            })
    return {
        "format": "K-Waay AROCMAG Stage 5 submission package manifest",
        "manifest_self_hash": "omitted to avoid recursive self-reference",
        "baseline_head": "b4b484f20f7b2670490725e8b0653e6ca6de360b",
        "status": "SUBMISSION_PACKAGE_PREPARED_WITH_BLOCKERS",
        "files": files,
    }


def main() -> None:
    required_paths = [PACKAGE / name for name in EXPECTED_PACKAGE_FILES]
    missing = [str(path.relative_to(ROOT)).replace("\\", "/") for path in required_paths if not path.is_file()]

    with zipfile.ZipFile(STAGE4_DOCX) as previous, zipfile.ZipFile(STAGE5_DOCX) as current:
        previous_names = previous.namelist()
        current_names = current.namelist()
        common_names = sorted(set(previous_names) & set(current_names))
        different_parts = [
            name for name in common_names
            if previous.read(name) != current.read(name)
        ]
        stage5_bad_member = current.testzip()
        document_xml = current.read("word/document.xml")
        document_root = etree.fromstring(document_xml)
        all_xml_text = "\n".join(
            current.read(name).decode("utf-8", "ignore")
            for name in current_names
            if name.endswith(".xml") or name.endswith(".rels")
        )
        comments_parts = [name for name in current_names if "comments" in name.lower() or name.endswith("people.xml")]
        tracked_counts = {
            "insertions": len(document_root.xpath(".//w:ins", namespaces=NS)),
            "deletions": len(document_root.xpath(".//w:del", namespaces=NS)),
            "moves_from": len(document_root.xpath(".//w:moveFrom", namespaces=NS)),
            "moves_to": len(document_root.xpath(".//w:moveTo", namespaces=NS)),
        }
        hidden_run_count = len(document_root.xpath(".//w:rPr/w:vanish | .//w:rPr/w:webHidden", namespaces=NS))
        external_relationship_parts = [
            name for name in current_names
            if name.endswith(".rels") and b'TargetMode="External"' in current.read(name)
        ]
        local_path_hits = sorted(set(re.findall(
            r"(?:[A-Za-z]:\\[^<\"']+|kwaay-formal|Users\\10202)",
            all_xml_text,
            re.I,
        )))
        omml_count = len(document_root.xpath(".//m:oMath", namespaces=NS))
        object_count = len(document_root.xpath(".//w:object", namespaces=NS))
        ole_count = len(document_root.xpath(".//o:OLEObject", namespaces=NS))
        embedding_parts = [name for name in current_names if name.startswith("word/embeddings/")]
        macro_parts = [name for name in current_names if "vbaProject" in name or name.endswith(".bin")]

    doc = Document(STAGE5_DOCX)
    core = doc.core_properties
    figure_records = []
    for shape in doc.inline_shapes:
        embed = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        part = doc.part.rels[embed].target_part
        data = part.blob
        image = Image.open(io.BytesIO(data))
        figure_records.append({
            "part": str(part.partname),
            "format": image.format,
            "pixels": list(image.size),
            "dpi": list(image.info.get("dpi", (None, None))),
            "sha256": sha256_bytes(data),
        })

    with FORMULA_CHECKLIST.open("r", encoding="utf-8-sig", newline="") as stream:
        formula_rows = list(csv.DictReader(stream))
    formula_indices = [int(row["object_index"]) for row in formula_rows]

    stage4_reader = PdfReader(str(STAGE4_PDF))
    stage5_reader = PdfReader(str(STAGE5_PDF))
    stage5_pdf_metadata = {str(key): str(value) for key, value in (stage5_reader.metadata or {}).items()}
    pdf_metadata_text = "\n".join(stage5_pdf_metadata.values())
    pdf_metadata_path_hits = sorted(set(re.findall(
        r"(?:[A-Za-z]:\\[^\r\n]+|kwaay-formal|Users\\10202)",
        pdf_metadata_text,
        re.I,
    )))
    stage5_page_sizes = [
        [float(page.mediabox.width), float(page.mediabox.height)]
        for page in stage5_reader.pages
    ]
    pdf_text_equal = normalized_pdf_text(STAGE4_PDF) == normalized_pdf_text(STAGE5_PDF)

    rendered_pages = []
    for number in range(1, 7):
        path = RENDER_DIR / f"verified-page-{number}.png"
        if path.is_file():
            with Image.open(path) as image:
                rendered_pages.append({
                    "page": number,
                    "pixels": list(image.size),
                    "sha256": sha256(path),
                })

    author_form = (PACKAGE / "AUTHOR_INFORMATION_COLLECTION_FORM.md").read_text(encoding="utf-8")
    email_hits = re.findall(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", author_form, re.I)
    phone_hits = re.findall(r"(?<!\d)1[3-9]\d{9}(?!\d)", author_form)

    checks = {
        "all_required_package_files_present": not missing,
        "stage5_docx_zip_integrity": stage5_bad_member is None,
        "same_docx_part_names_as_stage4": previous_names == current_names,
        "only_core_metadata_differs_from_stage4": different_parts == ["docProps/core.xml"],
        "document_xml_byte_identical_to_stage4": "word/document.xml" not in different_parts,
        "anonymous_core_metadata": core.author == "Anonymous Author(s)" and core.last_modified_by == "Anonymous Author(s)",
        "stage5_subject_present": core.subject == "Anonymous AROCMAG Stage 5 author-confirmation manuscript",
        "no_comments_or_people_parts": not comments_parts,
        "no_tracked_changes": not any(tracked_counts.values()),
        "no_hidden_runs": hidden_run_count == 0,
        "no_external_relationships": not external_relationship_parts,
        "no_local_paths": not local_path_hits,
        "no_macros": not macro_parts,
        "formula_count_is_99": omml_count == 99,
        "no_ole_or_embedded_objects": object_count == 0 and ole_count == 0 and not embedding_parts,
        "four_tables": len(doc.tables) == 4,
        "two_inline_figures": len(figure_records) == 2,
        "frozen_figure_hashes_preserved": sorted(item["sha256"] for item in figure_records) == sorted(EXPECTED_FIGURE_HASHES),
        "formula_checklist_has_99_rows": len(formula_rows) == 99,
        "formula_checklist_indices_are_1_to_99": formula_indices == list(range(1, 100)),
        "stage5_pdf_has_six_pages": len(stage5_reader.pages) == 6,
        "stage4_pdf_has_six_pages": len(stage4_reader.pages) == 6,
        "stage5_pdf_is_a4": all(abs(width - 595.28) < 1 and abs(height - 841.89) < 1 for width, height in stage5_page_sizes),
        "stage5_pdf_text_matches_stage4": pdf_text_equal,
        "stage5_pdf_author_is_anonymous": stage5_pdf_metadata.get("/Author") == "Anonymous Author(s)",
        "stage5_pdf_metadata_has_no_local_paths": not pdf_metadata_path_hits,
        "six_rendered_pages_present": len(rendered_pages) == 6,
        "rendered_pages_have_expected_size": all(item["pixels"] == [1241, 1754] for item in rendered_pages),
        "author_form_contains_no_filled_email_or_phone": not email_hits and not phone_hits,
    }

    manifest = package_manifest()
    MANIFEST_OUT.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    result = {
        "baseline_head": "b4b484f20f7b2670490725e8b0653e6ca6de360b",
        "stage4_docx_sha256": sha256(STAGE4_DOCX),
        "stage5_docx_sha256": sha256(STAGE5_DOCX),
        "stage4_pdf_sha256": sha256(STAGE4_PDF),
        "stage5_pdf_sha256": sha256(STAGE5_PDF),
        "docx_different_parts": different_parts,
        "docx_structure": {
            "omml": omml_count,
            "word_objects": object_count,
            "ole_objects": ole_count,
            "embedding_parts": embedding_parts,
            "tables": len(doc.tables),
            "inline_figures": len(figure_records),
        },
        "core_properties": {
            "author": core.author,
            "last_modified_by": core.last_modified_by,
            "subject": core.subject,
            "description": core.comments,
        },
        "privacy": {
            "comments_or_people_parts": comments_parts,
            "tracked_change_counts": tracked_counts,
            "hidden_run_count": hidden_run_count,
            "external_relationship_parts": external_relationship_parts,
            "local_path_hits": local_path_hits,
            "macro_parts": macro_parts,
            "author_form_email_hits": email_hits,
            "author_form_phone_hits": phone_hits,
        },
        "figures": figure_records,
        "formula_checklist_rows": len(formula_rows),
        "pdf_page_sizes_points": stage5_page_sizes,
        "pdf_metadata": stage5_pdf_metadata,
        "pdf_metadata_local_path_hits": pdf_metadata_path_hits,
        "rendered_pages": rendered_pages,
        "missing_package_files": missing,
        "checks": checks,
        "all_internal_checks_pass": all(checks.values()),
        "manual_visual_review": {
            "status": "completed separately",
            "pages_reviewed": [1, 2, 3, 4, 5, 6],
            "scope": "WPS-exported Stage 5 PDF at 150 dpi; not a Microsoft Word validation",
        },
        "blocked_statuses": [
            "MATHTYPE_CONVERSION_BLOCKED",
            "VECTOR_GRAPHICS_PENDING_IF_SUBMISSION_SYSTEM_REQUIRES_VECTOR",
            "MICROSOFT_WORD_VALIDATION_PENDING",
            "AUTHOR_METADATA_PENDING",
            "SUBMISSION_SYSTEM_RULES_PENDING",
        ],
        "final_status": "SUBMISSION_PACKAGE_PREPARED_WITH_BLOCKERS",
    }
    AUDIT_OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "audit": str(AUDIT_OUT),
        "manifest": str(MANIFEST_OUT),
        "checks": checks,
        "all_internal_checks_pass": result["all_internal_checks_pass"],
    }, ensure_ascii=False, indent=2))
    if not result["all_internal_checks_pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
