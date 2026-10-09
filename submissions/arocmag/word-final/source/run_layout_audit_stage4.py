from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[4]
WORD_ROOT = ROOT / "submissions" / "arocmag" / "word-final"
DOCX = WORD_ROOT / "K-Waay-AROCMAG-submission-stage4.docx"
PDF = WORD_ROOT / "qa" / "stage4-render-final" / "K-Waay-AROCMAG-submission-stage4-verified.pdf"
OUT = WORD_ROOT / "qa" / "stage4-layout-audit.json"


def twips_to_cm(value: int) -> float:
    return round(value / 1440 * 2.54, 3)


def section_column_count(section) -> int:
    cols = section._sectPr.xpath("./w:cols")
    if not cols:
        return 1
    return int(cols[0].get(qn("w:num"), "1"))


def main() -> None:
    doc = Document(DOCX)
    with zipfile.ZipFile(DOCX) as archive:
        bad_zip_member = archive.testzip()
        names = archive.namelist()
        xml = archive.read("word/document.xml")
        media = sorted(name for name in names if name.startswith("word/media/"))
        external_relationships = []
        for name in names:
            if name.endswith(".rels"):
                text = archive.read(name).decode("utf-8", "ignore")
                if 'TargetMode="External"' in text:
                    external_relationships.append(name)

    sections = []
    for section in doc.sections:
        sections.append({
            "page_width_cm": round(section.page_width.cm, 3),
            "page_height_cm": round(section.page_height.cm, 3),
            "top_margin_cm": round(section.top_margin.cm, 3),
            "bottom_margin_cm": round(section.bottom_margin.cm, 3),
            "left_margin_cm": round(section.left_margin.cm, 3),
            "right_margin_cm": round(section.right_margin.cm, 3),
            "columns": section_column_count(section),
        })

    image_parts = []
    image_widths_cm = []
    for shape in doc.inline_shapes:
        embed = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        rel = doc.part.rels[embed]
        image_parts.append(str(rel.target_part.partname))
        image_widths_cm.append(round(shape.width.cm, 3))

    table_shapes = [[len(table.rows), len(table.columns)] for table in doc.tables]
    display_formula_count = sum(p.style.style_id == "74" for p in doc.paragraphs)
    omath_count = len(doc.element.xpath(".//*[local-name()='oMath']"))
    column_break_count = len(doc.element.xpath(".//*[local-name()='br' and @*[local-name()='type']='column']"))
    ole_object_count = len(doc.element.xpath(".//*[local-name()='OLEObject']"))
    embedded_object_parts = [name for name in names if name.startswith("word/embeddings/")]

    reader = PdfReader(str(PDF))
    first_box = reader.pages[0].mediabox
    pdf_width_pt = float(first_box.width)
    pdf_height_pt = float(first_box.height)

    checks = {
        "zip_integrity": bad_zip_member is None,
        "a4_all_sections": all(
            abs(s["page_width_cm"] - 21.0) < 0.02 and abs(s["page_height_cm"] - 29.7) < 0.02
            for s in sections
        ),
        "official_margins_all_sections": all(
            abs(s["top_margin_cm"] - 2.0) < 0.02
            and abs(s["bottom_margin_cm"] - 1.5) < 0.02
            and abs(s["left_margin_cm"] - 1.5) < 0.02
            and abs(s["right_margin_cm"] - 1.5) < 0.02
            for s in sections
        ),
        "single_column_front_two_column_body": sections[0]["columns"] == 1 and sections[1]["columns"] == 2,
        "full_width_table_sections_present": sum(s["columns"] == 1 for s in sections) == 5,
        "four_tables_exact_shapes": table_shapes == [[6, 5], [4, 5], [5, 3], [8, 5]],
        "two_single_column_figures": len(image_parts) == 2 and all(abs(width - 7.9) < 0.02 for width in image_widths_cm),
        "paper_figures_are_png": all(part.lower().endswith(".png") for part in image_parts),
        "six_display_formulas": display_formula_count == 6,
        "editable_omml_present": omath_count == 99,
        "no_ole_or_mathtype_objects": ole_object_count == 0 and not embedded_object_parts,
        "one_final_column_break": column_break_count == 1,
        "no_external_package_relationships": not external_relationships,
        "pdf_six_pages": len(reader.pages) == 6,
        "pdf_a4": abs(pdf_width_pt - 595.3) < 1.0 and abs(pdf_height_pt - 841.9) < 1.0,
        "anonymous_core_metadata": (
            doc.core_properties.author == "Anonymous Author(s)"
            and doc.core_properties.last_modified_by == "Anonymous Author(s)"
        ),
    }

    result = {
        "checks": checks,
        "all_pass": all(checks.values()),
        "sections": sections,
        "table_shapes": table_shapes,
        "image_parts": image_parts,
        "image_widths_cm": image_widths_cm,
        "package_media": media,
        "display_formula_count": display_formula_count,
        "omath_count": omath_count,
        "ole_object_count": ole_object_count,
        "embedded_object_parts": embedded_object_parts,
        "column_break_count": column_break_count,
        "external_relationship_parts": external_relationships,
        "pdf_pages": len(reader.pages),
        "pdf_page_size_pt": [round(pdf_width_pt, 2), round(pdf_height_pt, 2)],
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["all_pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
