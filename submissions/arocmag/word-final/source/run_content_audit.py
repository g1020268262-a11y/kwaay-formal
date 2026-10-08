from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[4]
WORD_ROOT = ROOT / "submissions" / "arocmag" / "word-final"
DOCX = WORD_ROOT / "K-Waay-AROCMAG-submission-working.docx"
OUT = WORD_ROOT / "qa" / "content-audit.json"


def norm(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def package_visible_text(path: Path) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name == "word/document.xml" or re.fullmatch(r"word/(header|footer)\d+\.xml", name):
                root = etree.fromstring(archive.read(name))
                chunks.extend(root.xpath("//*[local-name()='t']/text()"))
    return "\n".join(chunks)


def main() -> None:
    doc = Document(DOCX)
    visible = package_visible_text(DOCX)
    compact = norm(visible)
    squash = re.sub(r"\s+", "", visible)
    paragraph_texts = [norm(paragraph.text) for paragraph in doc.paragraphs]

    expected_zh = (
        "K-Waay的BatchReceive 已要求同次调用输入对应不同参与方，但批内消息互异性以及发送来源与接受事件之间的单射关系能否替代该批内参与方关系仍需明确。"
        "本文在统一的发送、来源匹配、两槽收集和顺序处理生命周期上，构造无互异、消息互异和参与方互异三种Tamarin批次准入模型。"
        "验证表明：无互异模型中，一个匹配发送来源可对应同批两个接受事件；消息互异模型的批内消息互异性和批内来源单射性均验证通过，但同一参与方的不同发送实例及消息仍可同批接受；"
        "参与方互异模型保持批内参与方互异性，且有效的不同参与方批次仍可达。同参与方异消息的可达执行表明，消息级约束不能推出批内参与方互异性。上述结论限于固定两槽的批次准入抽象。"
    )
    expected_en = (
        "K-Waay already requires the inputs of one BatchReceive invocation to correspond to different parties. "
        "It remains necessary to determine whether message distinction and an injective correspondence between sender origins and acceptance events can substitute for this same-batch party relation. "
        "We construct three fixed-two-slot Tamarin models that share the same lifecycle for sending, origin matching, collection, admission, and sequential processing. "
        "The models respectively impose no distinction constraint, message distinction, and party distinction. In the first model, one matching sender origin can correspond to two acceptance events in the same batch. "
        "In the message-constrained model, same-batch message distinction and batch-scoped origin injectivity are both verified, yet two different sender occurrences and messages from the same party can still be accepted in one batch. "
        "The party-constrained model preserves same-batch party distinction and still admits a valid batch formed by different parties. "
        "The reachable same-party, different-message execution shows that the message-level restriction does not imply same-batch party distinction. The conclusion is limited to the fixed-two-slot batch-admission abstraction."
    )
    zh_para = next(text for text in paragraph_texts if text.startswith("摘 要：") or text.startswith("摘 要：") or text.startswith("摘  要："))
    en_para = next(text for text in paragraph_texts if text.startswith("Abstract:"))
    zh_actual = re.sub(r"^摘\s*要：", "", zh_para).strip()
    en_actual = en_para.removeprefix("Abstract:").strip()

    forbidden_patterns = {
        "latex_cite": r"\\cite(?:\[|\{)",
        "latex_ref": r"\\(?:eqref|ref)\{",
        "latex_environment": r"\\(?:begin|end|input)\{",
        "latex_project_macros": r"\\(?:batchreceive|receiveraccept|sendaction|oid|kwaay)\b",
        "todo": r"\bTODO\b",
        "local_path": r"(?:[A-Za-z]:\\|kwaay-formal|word-draft)",
        "github_internal": r"github\.com/.+/(?:blob|tree)/",
    }
    forbidden_hits = {name: re.findall(pattern, visible, flags=re.I) for name, pattern in forbidden_patterns.items()}

    anchors = [
        "same_party_different_messages_batch_exists",
        "全部14项结果合计13项验证通过、1项被反例否定",
        "批内参与方互异性蕴含批内消息互异性",
        "不是新增的机器证明",
        "机器证据固定为两个槽位",
        "不能自动解释为具体协议攻击",
        "本文结论限于两槽准入抽象及ReceiverAccept事件",
        "该元组的匹配Send唯一，执行中可能存在与其不匹配的其他Send",
    ]
    anchor_status = {anchor: re.sub(r"\s+", "", anchor) in squash for anchor in anchors}

    table_shapes = [[len(table.rows), len(table.columns)] for table in doc.tables]
    heading_counts = {
        "level1": sum(paragraph.style.style_id == "65" for paragraph in doc.paragraphs),
        "level2": sum(paragraph.style.style_id == "66" for paragraph in doc.paragraphs),
    }
    omath_count = len(doc.element.xpath(".//*[local-name()='oMath']"))
    inline_shapes = len(doc.inline_shapes)
    numbered_equations = [number for number in (1, 2, 3) if any(f"({number})" in text for text in paragraph_texts)]
    cited = set()
    for group in re.findall(r"\[([0-9,，\s]+)\]", compact):
        cited.update(int(value) for value in re.findall(r"\d+", group))
    citation_numbers = sorted(number for number in cited if 1 <= number <= 8)

    core = doc.core_properties
    metadata = {
        "title": core.title,
        "subject": core.subject,
        "author": core.author,
        "last_modified_by": core.last_modified_by,
    }
    checks = {
        "title_exact": TITLE_ZH == paragraph_texts[0],
        "chinese_abstract_exact": norm(expected_zh) == norm(zh_actual),
        "english_abstract_exact": norm(expected_en) == norm(en_actual),
        "keywords_present": "关键词：K-Waay；批次准入；批内参与方互异条件；Tamarin；形式化验证" in squash,
        "english_keywords_present": "Key words: K-Waay; batch admission; distinct-party condition; Tamarin; formal verification" in compact,
        "all_scientific_anchors": all(anchor_status.values()),
        "four_tables_exact_shapes": table_shapes == [[6, 5], [4, 5], [5, 3], [8, 5]],
        "two_figures": inline_shapes == 2,
        "three_numbered_equations": numbered_equations == [1, 2, 3],
        "editable_omml_present": omath_count > 0,
        "all_eight_citations_present": citation_numbers == list(range(1, 9)),
        "heading_counts": heading_counts == {"level1": 6, "level2": 13},
        "no_forbidden_visible_text": all(not values for values in forbidden_hits.values()),
        "anonymous_metadata": metadata["author"] == "Anonymous Author(s)" and metadata["last_modified_by"] == "Anonymous Author(s)",
    }
    result = {
        "checks": checks,
        "all_pass": all(checks.values()),
        "anchor_status": anchor_status,
        "forbidden_hits": forbidden_hits,
        "table_shapes": table_shapes,
        "heading_counts": heading_counts,
        "omath_count": omath_count,
        "inline_shapes": inline_shapes,
        "numbered_equations": numbered_equations,
        "citation_numbers": citation_numbers,
        "metadata": metadata,
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["all_pass"]:
        raise SystemExit(1)


TITLE_ZH = "K-Waay批次准入中参与方互异条件的形式化分析"


if __name__ == "__main__":
    main()
