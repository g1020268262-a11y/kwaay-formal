from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[4]
OUT = ROOT / "submissions" / "arocmag" / "word-draft"
SOURCE, MANUSCRIPT = OUT / "source", OUT / "manuscript"
REPORTS, REFOUT = OUT / "reports", OUT / "references"
SECTIONS = ROOT / "submissions" / "arocmag" / "draft" / "sections"
BIB_PATH = ROOT / "manuscript" / "bibliography" / "references.bib"
BASE_DOCX = SOURCE / "official-template-converted.docx"
FINAL_DOCX = MANUSCRIPT / "kwaay-arocmag-draft.docx"
ACCESS_DATE = "2026-10-06"
FIGURE_MARKERS = {1: "[[FIGURE_1]]", 2: "[[FIGURE_2]]"}


def ensure_dirs():
    for p in (SOURCE, MANUSCRIPT, REPORTS, REFOUT):
        p.mkdir(parents=True, exist_ok=True)
    copied = SOURCE / "manuscript-sections"
    copied.mkdir(parents=True, exist_ok=True)
    for src in sorted(SECTIONS.glob("*.md")):
        shutil.copy2(src, copied / src.name)
    shutil.copy2(BIB_PATH, SOURCE / "references.bib")


def parse_bib(text):
    entries, pos = {}, 0
    while True:
        m = re.search(r"@(\w+)\s*\{\s*([^,\s]+)\s*,", text[pos:], re.S)
        if not m:
            break
        kind, key, start = m.group(1).lower(), m.group(2), pos + m.end()
        depth, i = 1, start
        while i < len(text) and depth:
            depth += (text[i] == "{") - (text[i] == "}")
            i += 1
        body, fields, fpos = text[start:i - 1], {"ENTRYTYPE": kind, "ID": key}, 0
        while fpos < len(body):
            fm = re.search(r"(\w+)\s*=\s*", body[fpos:])
            if not fm:
                break
            name, j = fm.group(1).lower(), fpos + fm.end()
            if body[j] == "{":
                dep, k = 1, j + 1
                while k < len(body) and dep:
                    dep += (body[k] == "{") - (body[k] == "}")
                    k += 1
                value, fpos = body[j + 1:k - 1], k
            elif body[j] == '"':
                k = body.find('"', j + 1)
                value, fpos = body[j + 1:k], k + 1
            else:
                k = body.find(",", j)
                k = len(body) if k < 0 else k
                value, fpos = body[j:k].strip(), k + 1
            fields[name] = re.sub(r"\s+", " ", value.strip())
        entries[key], pos = fields, i
    return entries


ACCENTS = {
    r'{\"i}': "ï", r'{\"e}': "ë", r'{\"o}': "ö", r'{\"u}': "ü",
    r"{\'e}": "é", r"{\'a}": "á", r"{\'i}": "í",
    r"{\v{s}}": "š", r"{\o}": "ø", r"{\O}": "Ø",
}


def unlatex(s):
    s = s.replace(r"\url", "")
    for old, new in ACCENTS.items():
        s = s.replace(old, new)
    s = re.sub(r"\{\\[\"'`^~=.uvHckbdtr]\s*([A-Za-z])\}", r"\1", s)
    return re.sub(r"\s+", " ", s.replace("{", "").replace("}", "").replace(r"\&", "&").replace("~", " ")).strip()


def author_list(text):
    people = [unlatex(x.strip()) for x in re.split(r"\s+and\s+", text)]
    return ", ".join(people[:3] + (["et al"] if len(people) > 3 else []))


def scan_citation_order(files):
    order = []
    for path in files:
        text = path.read_text(encoding="utf-8")
        for m in re.finditer(r"\\cite(?:\[[^\]]+\])?\{([^}]+)\}", text):
            for key in m.group(1).split(","):
                key = key.strip()
                if key and key not in order:
                    order.append(key)
    return order


def format_reference(n, key, entry):
    authors, title = author_list(entry.get("author", "")), unlatex(entry.get("title", ""))
    year, pages = entry.get("year", ""), entry.get("pages", "").replace("--", "-")
    kind, needs = entry.get("ENTRYTYPE", ""), False
    if kind == "article":
        journal = unlatex(entry.get("journal", ""))
        tail = f"{journal}, {year}"
        if entry.get("volume"):
            tail += f", {entry['volume']}"
            if entry.get("number"):
                tail += f"({entry['number']})"
        if pages:
            tail += f": {pages}"
        text = f"[{n}] {authors}. {title}[J]. {tail}."
    elif kind == "inproceedings":
        tail = f"{unlatex(entry.get('booktitle', ''))}, {year}"
        if pages:
            tail += f": {pages}"
        text, needs = f"[{n}] {authors}. {title}[C]//{tail}.", True
    elif kind == "techreport":
        text = f"[{n}] {authors}. {title}[R/OL]. {unlatex(entry.get('institution',''))}, {unlatex(entry.get('number',''))}, {year}"
        if entry.get("url"):
            text += f"[{ACCESS_DATE}]. {entry['url']}"
        text, needs = text + ".", True
    else:
        text = f"[{n}] {authors}. {title}[EB/OL]. {unlatex(entry.get('howpublished',''))}, {year}"
        if entry.get("url"):
            text += f"[{ACCESS_DATE}]. {entry['url']}"
        text, needs = text + ".", True
    return re.sub(r"\s+", " ", text), needs


def style_by_id(doc, sid):
    return next(s for s in doc.styles if s.style_id == sid)


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=9, bold=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size, run.font.color.rgb = Pt(size), RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


SUB = str.maketrans("0123456789+-=()aeioruvxlmn", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑᵢₒᵣᵤᵥₓₗₘₙ")
SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
COMMANDS = {
    "ne": "≠", "neq": "≠", "ge": "≥", "geq": "≥", "le": "≤", "leq": "≤",
    "land": "∧", "lor": "∨", "Rightarrow": "⇒", "Leftrightarrow": "⇔",
    "leftarrow": "←", "rightarrow": "→", "triangleq": "≜", "equiv": "≡",
    "forall": "∀", "exists": "∃", "in": "∈", "notin": "∉", "bot": "⊥",
    "tau": "τ", "ell": "ℓ", "cdot": "·", "ldots": "…", "dots": "…",
    "langle": "⟨", "rangle": "⟩", "times": "×", "mapsto": "↦",
    "longrightarrow": "→", "Longrightarrow": "⇒",
}


def extract_group(s, start):
    if start >= len(s):
        return "", start
    if s[start] != "{":
        return s[start], start + 1
    dep, i = 1, start + 1
    while i < len(s) and dep:
        dep += (s[i] == "{") - (s[i] == "}")
        i += 1
    return s[start + 1:i - 1], i


def replace_scripts(s):
    for marker, table in (("_", SUB), ("^", SUP)):
        i, pieces = 0, []
        while i < len(s):
            if s[i] == marker and i + 1 < len(s):
                group, end = extract_group(s, i + 1)
                pieces.append(latex_to_unicode(group, scripts=False).translate(table))
                i = end
            else:
                pieces.append(s[i])
                i += 1
        s = "".join(pieces)
    return s


def latex_to_unicode(src, scripts=True):
    s = src.strip().replace(r"\not\Rightarrow", "⇏")
    s = s.replace(r"\{", "⦃").replace(r"\}", "⦄")
    s = re.sub(r"\\begin\{(?:aligned|array)\}(?:\{[^}]*\})?", "", s)
    s = re.sub(r"\\end\{(?:aligned|array)\}", "", s)
    s = s.replace("&", "").replace(r"\\", "\n")
    for cmd, symbol in sorted(COMMANDS.items(), key=lambda x: -len(x[0])):
        s = re.sub(r"\\" + re.escape(cmd) + r"\b", symbol, s)
    s = re.sub(r"\\(?:left|right|bigl|bigr|Bigl|Bigr)\b", "", s)
    for cmd in ("operatorname", "mathrm", "mathsf", "texttt", "text", "mathbf", "mathit"):
        for _ in range(5):
            s = re.sub(r"\\" + cmd + r"\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", lambda m: f"({m.group(1)})/({m.group(2)})", s)
    s = s.replace(r"\,", " ").replace(r"\;", " ").replace(r"\:", " ")
    s = s.replace(r"\quad", "  ").replace(r"\qquad", "    ")
    s = re.sub(r"\\([A-Za-z]+)", r"\1", s)
    s = re.sub(r"\\(?=\s|$)", "", s)
    s = replace_scripts(s) if scripts else s
    s = re.sub(r"[ \t]+", " ", s.replace("{", "").replace("}", ""))
    return re.sub(r" *\n *", "\n", s).strip()


def omath(text):
    o, r = OxmlElement("m:oMath"), OxmlElement("m:r")
    rpr, sty = OxmlElement("m:rPr"), OxmlElement("m:sty")
    sty.set(qn("m:val"), "p")
    rpr.append(sty)
    r.append(rpr)
    t = OxmlElement("m:t")
    t.set(qn("xml:space"), "preserve")
    t.text = latex_to_unicode(text)
    r.append(t)
    o.append(r)
    return o


def add_text_run(p, text, bold=False, italic=False, size=9):
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold)
        r.italic = italic


def replace_citations(text, cite_map):
    def repl(m):
        detail = m.group(1)
        nums = [cite_map[k.strip()] for k in m.group(2).split(",") if k.strip() in cite_map]
        label = "[" + ",".join(str(x) for x in nums) + "]"
        return "（见文献" + label + "，" + detail.replace("--", "–") + "）" if detail else label
    return re.sub(r"\\cite(?:\[([^\]]+)\])?\{([^}]+)\}", repl, text)


def add_inline(p, text, cite_map, size=9):
    text = replace_citations(text, cite_map).replace(chr(96), "").replace("**", "")
    pos = 0
    for m in re.finditer(r"\$([^$]+)\$", text):
        add_text_run(p, text[pos:m.start()], size=size)
        p._p.append(omath(m.group(1)))
        pos = m.end()
    add_text_run(p, text[pos:], size=size)


def set_paragraph_metrics(p, first_indent=True, before=0, after=0, line=12):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(line)
    if first_indent:
        pf.first_line_indent = Pt(18)


def clear_document_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_columns(section, count, space_twips=360):
    sectPr, cols = section._sectPr, section._sectPr.xpath("./w:cols")
    el = cols[0] if cols else OxmlElement("w:cols")
    el.set(qn("w:num"), str(count))
    el.set(qn("w:space"), str(space_twips))
    if not cols:
        sectPr.append(el)


def set_section_geometry(section):
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(1.5)
    section.left_margin, section.right_margin = Cm(1.5), Cm(1.5)
    section.header_distance, section.footer_distance = Cm(0.7), Cm(0.7)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=style_by_id(doc, "65" if level == 1 else "66"))
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(5 if level == 1 else 3), Pt(1)
    add_text_run(p, text, bold=True, size=10.5 if level == 1 else 9)
    return p


def add_body_paragraph(doc, text, cite_map, indent=True):
    p = doc.add_paragraph(style=style_by_id(doc, "71"))
    set_paragraph_metrics(p, first_indent=indent)
    add_inline(p, text, cite_map)
    return p


def add_display_math(doc, lines, number):
    p = doc.add_paragraph(style=style_by_id(doc, "74"))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(2), Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.3), WD_TAB_ALIGNMENT.RIGHT)
    for idx, line in enumerate(lines):
        if not line.strip():
            continue
        if idx:
            p.add_run().add_break()
        p._p.append(omath(line))
    r = p.add_run("\t(" + str(number) + ")")
    set_run_font(r, size=9)
    return p


def set_cell_text(cell, text, cite_map, bold=False, size=7.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(10)
    add_inline(p, text, cite_map, size=size)
    for r in p.runs:
        set_run_font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, attrs in edges.items():
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        for key, val in attrs.items():
            el.set(qn("w:" + key), str(val))


def add_table(doc, rows, cite_map):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, True
    none = {"val": "nil"}
    rule = {"val": "single", "sz": "6", "space": "0", "color": "000000"}
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            set_cell_text(cell, value, cite_map, bold=(i == 0), size=7 if len(row) >= 5 else 7.5)
            set_cell_border(cell, top=rule if i == 0 else none,
                            bottom=rule if i in (0, len(rows) - 1) else none,
                            left=none, right=none, insideH=none, insideV=none)
    return table


def parse_md_table(lines):
    rows = []
    for line in lines:
        vals = [x.strip() for x in line.strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", x) for x in vals):
            rows.append(vals)
    return rows


def add_caption(doc, cn, en=""):
    p = doc.add_paragraph(style=style_by_id(doc, "77"))
    p.alignment, p.paragraph_format.keep_with_next = WD_ALIGN_PARAGRAPH.CENTER, True
    add_text_run(p, cn, bold=True, size=8)
    if en:
        p.add_run().add_break()
        add_text_run(p, en, size=8)
    return p


def add_front_matter(doc):
    text = (SECTIONS / "abstract-zh-en.md").read_text(encoding="utf-8")
    zh_title = re.search(r"# 中文题目\s+(.+?)\s+## 中文摘要", text, re.S).group(1).strip()
    zh_abs = re.search(r"## 中文摘要\s+(.+?)\s+\*\*关键词：\*\*", text, re.S).group(1).strip()
    zh_kw = re.search(r"\*\*关键词：\*\*\s*(.+)", text).group(1).strip()
    en_title = re.search(r"# English Title\s+(.+?)\s+## Abstract", text, re.S).group(1).strip()
    en_abs = re.search(r"## Abstract\s+(.+?)\s+\*\*Key words:\*\*", text, re.S).group(1).strip()
    en_kw = re.search(r"\*\*Key words:\*\*\s*(.+)", text).group(1).strip()

    p = doc.add_paragraph(style=style_by_id(doc, "40"))
    p.alignment, p.paragraph_format.space_after = WD_ALIGN_PARAGRAPH.CENTER, Pt(4)
    add_text_run(p, zh_title, bold=True, size=16)
    for line, size in (
        ("作者A，作者B（待补充并核验）", 11),
        ("（作者单位、城市、邮编待补充并核验）", 9),
        ("基金项目、作者简介及通信作者信息待补充并核验", 8),
    ):
        p = doc.add_paragraph(style=style_by_id(doc, "41"))
        p.alignment, p.paragraph_format.space_after = WD_ALIGN_PARAGRAPH.CENTER, Pt(1)
        add_text_run(p, line, size=size)
    p = doc.add_paragraph(style=style_by_id(doc, "45"))
    set_paragraph_metrics(p, first_indent=False)
    add_text_run(p, "摘  要：", bold=True)
    add_text_run(p, zh_abs)
    p = doc.add_paragraph(style=style_by_id(doc, "45"))
    set_paragraph_metrics(p, first_indent=False)
    add_text_run(p, "关键词：", bold=True)
    add_text_run(p, zh_kw)
    p = doc.add_paragraph(style=style_by_id(doc, "42"))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(5), Pt(2)
    add_text_run(p, en_title, bold=True, size=12)
    p = doc.add_paragraph(style=style_by_id(doc, "56"))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text_run(p, "Author A, Author B (to be completed and verified)", size=9)
    p = doc.add_paragraph(style=style_by_id(doc, "45"))
    set_paragraph_metrics(p, first_indent=False, line=11)
    add_text_run(p, "Abstract: ", bold=True, size=8.5)
    add_text_run(p, en_abs, size=8.5)
    p = doc.add_paragraph(style=style_by_id(doc, "45"))
    set_paragraph_metrics(p, first_indent=False, line=11)
    add_text_run(p, "Key words: ", bold=True, size=8.5)
    add_text_run(p, en_kw, size=8.5)


def add_coordinate_legend(doc, rows, cite_map):
    p = doc.add_paragraph(style=style_by_id(doc, "71"))
    set_paragraph_metrics(p, first_indent=False)
    add_text_run(p, "坐标说明：", bold=True)
    for row in rows[1:]:
        if len(row) >= 3:
            before = len(p.runs)
            add_inline(p, row[0] + "——", cite_map)
            for run in p.runs[before:]:
                run.bold = True
            add_inline(p, row[1] + "；" + row[2] + "。", cite_map)


def strip_bold(line):
    return line.replace("**", "").strip()


def process_body(doc, files, cite_map):
    eq_no, formula_records, pending_table_caption = 0, [], None
    for path in files:
        lines, i = path.read_text(encoding="utf-8").splitlines(), 0
        while i < len(lines):
            line = lines[i].rstrip()
            if not line:
                i += 1
                continue
            if line.startswith("# "):
                heading = re.sub(r"^\d+(?:\.\d+)*\s+", "", line[2:].strip())
                add_heading(doc, heading, 1)
                i += 1
                continue
            if line.startswith("## "):
                heading = re.sub(r"^\d+(?:\.\d+)*\s+", "", line[3:].strip())
                add_heading(doc, heading, 2)
                i += 1
                continue
            if line.strip() == "$$":
                i += 1
                math_lines = []
                while i < len(lines) and lines[i].strip() != "$$":
                    math_lines.append(lines[i])
                    i += 1
                eq_no += 1
                raw = "\n".join(math_lines).strip()
                split = [x.strip() for x in re.split(r"\\\\", raw) if x.strip()]
                add_display_math(doc, split or [raw], eq_no)
                formula_records.append({"number": eq_no, "source_file": path.name,
                                        "latex": raw, "rendered": latex_to_unicode(raw)})
                i += 1
                continue
            if line.startswith("> **图"):
                cn, en = strip_bold(line[2:]), ""
                i += 1
                if i < len(lines) and lines[i].startswith("> **Fig."):
                    en = strip_bold(lines[i][2:])
                    i += 1
                note_lines = []
                while i < len(lines) and (not lines[i].strip() or lines[i].startswith(">")):
                    if lines[i].startswith(">") and "注：" in lines[i]:
                        note_lines.append(lines[i].lstrip("> ").strip())
                    i += 1
                number = int(re.search(r"图(\d+)", cn).group(1))
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(2), Pt(1)
                add_text_run(p, FIGURE_MARKERS[number], size=8)
                add_caption(doc, cn, en)
                p = doc.add_paragraph(style=style_by_id(doc, "80"))
                p.paragraph_format.space_after = Pt(2)
                add_text_run(p, " ".join(note_lines) or
                             "注：依据论文源材料人工整理；符号模型结论不扩展为完整协议精化或部署攻击。",
                             size=7.5)
                continue
            if re.match(r"\*\*表[123]\s", line):
                cn, en = strip_bold(line), ""
                i += 1
                if i < len(lines) and re.match(r"\*\*Table\s+[123]", lines[i]):
                    en = strip_bold(lines[i])
                    i += 1
                pending_table_caption = (cn, en)
                continue
            if line.startswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].startswith("|"):
                    table_lines.append(lines[i])
                    i += 1
                rows = parse_md_table(table_lines)
                if pending_table_caption:
                    sec = doc.add_section(WD_SECTION.CONTINUOUS)
                    set_section_geometry(sec)
                    set_columns(sec, 1)
                    add_caption(doc, *pending_table_caption)
                    add_table(doc, rows, cite_map)
                    sec = doc.add_section(WD_SECTION.CONTINUOUS)
                    set_section_geometry(sec)
                    set_columns(sec, 2)
                    pending_table_caption = None
                else:
                    add_coordinate_legend(doc, rows, cite_map)
                continue
            if line.startswith(">"):
                p = doc.add_paragraph(style=style_by_id(doc, "71"))
                set_paragraph_metrics(p, first_indent=False)
                p.paragraph_format.left_indent = Pt(10)
                add_inline(p, line.lstrip("> ").strip(), cite_map)
                i += 1
                continue
            para = [line]
            i += 1
            while i < len(lines):
                nxt = lines[i].rstrip()
                if (not nxt or nxt.startswith("#") or nxt.strip() == "$$"
                        or nxt.startswith("|") or nxt.startswith(">")
                        or re.match(r"\*\*表[123]\s", nxt)):
                    break
                para.append(nxt)
                i += 1
            text = " ".join(x.strip() for x in para)
            if text.startswith("**") and text.endswith("**"):
                p = add_body_paragraph(doc, strip_bold(text), cite_map, indent=False)
                for run in p.runs:
                    run.bold = True
            else:
                add_body_paragraph(doc, text, cite_map)
    (SOURCE / "formulas.json").write_text(
        json.dumps(formula_records, ensure_ascii=False, indent=2), encoding="utf-8")
    return eq_no, formula_records


def add_references(doc, order, bib):
    p = doc.add_paragraph(style=style_by_id(doc, "82"))
    p.paragraph_format.keep_with_next = True
    add_text_run(p, "参考文献", bold=True, size=9)
    records = []
    for n, key in enumerate(order, 1):
        text, needs = format_reference(n, key, bib[key])
        p = doc.add_paragraph(style=style_by_id(doc, "84"))
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Pt(18), Pt(-18)
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(9)
        display_text = re.sub(r"^\[\d+\]\s*", "", text)
        add_text_run(p, display_text, size=7.5)
        records.append({"number": n, "key": key, "text": text, "needs_confirmation": needs})
    (REFOUT / "citation-order.json").write_text(
        json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    with (REFOUT / "citation-order.tsv").open("w", encoding="utf-8", newline="\n") as f:
        f.write("number\tkey\tneeds_confirmation\treference\n")
        for r in records:
            f.write(f"{r['number']}\t{r['key']}\t{str(r['needs_confirmation']).lower()}\t{r['text']}\n")
    return records


def set_core_properties(doc):
    cp = doc.core_properties
    cp.title = "K-Waay批处理接纳中参与方区分条件的形式化分析"
    cp.subject, cp.author = "《计算机应用研究》投稿初稿", "待补充并核验"
    cp.keywords = "K-Waay; BatchReceive; 形式化验证; Tamarin; 批处理接纳"
    cp.comments = "由官方投稿模板转换并生成；作者与基金信息为待补充占位符。"


def normalize_styles(doc):
    for sid, east, latin, size in (
        ("65", "黑体", "Arial", 10.5), ("66", "黑体", "Arial", 9),
        ("71", "宋体", "Times New Roman", 9), ("74", "宋体", "Cambria Math", 9),
        ("77", "宋体", "Times New Roman", 8), ("80", "宋体", "Times New Roman", 7.5),
        ("82", "黑体", "Arial", 9), ("84", "宋体", "Times New Roman", 7.5),
    ):
        style = style_by_id(doc, sid)
        style.font.name, style.font.size = latin, Pt(size)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)


def write_reports(eq_count, formulas, refs, order):
    inline_count = sum(
        len(re.findall(r"(?<!\$)\$[^$\n]+\$(?!\$)", p.read_text(encoding="utf-8")))
        for p in sorted(SECTIONS.glob("*.md"))
    )
    formula_report = f"""# 公式审计

- 行间公式块：{eq_count} 个，已编号为（1）—（{eq_count}）。
- 行内公式源片段：{inline_count} 个。
- 写入格式：Office Math Markup Language（OMML），可在 Word/WPS 中编辑。
- 原始 LaTeX：`source/formulas.json`。
- MathType 状态：当前环境未发现可自动批量转换并复核的 MathType 接口，尚未完成 MathType 对象转换。
- 交付判定影响：按照任务约束，最终状态必须为 `AROCMAG_WORD_DRAFT_INCOMPLETE`。

可编辑性审计采用 OOXML 中的 `m:oMath` 元素计数，并在最终 WPS 打开后复核 OMath 数量。OMML 中的数学内容使用期刊可读的 Unicode 数学符号；复杂多行式保持可编辑，但后续如获得 MathType 环境，仍需逐式转换和复核。
"""
    (REPORTS / "FORMULA_AUDIT.md").write_text(formula_report, encoding="utf-8")
    pending = [r for r in refs if r["needs_confirmation"]]
    ref_report = f"""# 参考文献审计

- 正文首次引用顺序：{len(order)} 条。
- 参考文献表条目：{len(refs)} 条。
- 顺序映射文件：`references/citation-order.tsv` 与 `references/citation-order.json`。
- 正文引文已由 BibTeX 键转换为顺序编码 `[n]`；标题、摘要、关键词中未放置引文。
- 作者超过 3 人的条目采用“前 3 名 + et al”。
- 在线资源访问日期统一记录为 {ACCESS_DATE}。

## 仍需人工核验

{len(pending)} 条会议论文、技术报告或在线规范缺少官方模板示例所需的完整出版地/出版者或载体细节。稿件未猜测这些字段，已在映射 JSON 中标记 `needs_confirmation=true`。这些待核验项不影响正文引文编号闭合，但投稿前应按期刊编辑部的 GB/T 7714 口径补全。
"""
    (REPORTS / "REFERENCE_AUDIT.md").write_text(ref_report, encoding="utf-8")


def main():
    ensure_dirs()
    files = [SECTIONS / x for x in (
        "00-introduction.md", "01-batch-admission-problem.md", "02-symbolic-modeling.md",
        "03-formal-analysis.md", "04-discussion.md", "05-related-work.md", "06-conclusion.md")]
    bib = parse_bib(BIB_PATH.read_text(encoding="utf-8"))
    order = scan_citation_order(files)
    missing = [k for k in order if k not in bib]
    if missing:
        raise RuntimeError("missing bibliography entries: " + ", ".join(missing))
    cite_map = {key: i for i, key in enumerate(order, 1)}
    doc = Document(BASE_DOCX)
    clear_document_body(doc)
    normalize_styles(doc)
    first = doc.sections[0]
    set_section_geometry(first)
    set_columns(first, 1)
    add_front_matter(doc)
    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_section_geometry(body_section)
    set_columns(body_section, 2)
    eq_count, formulas = process_body(doc, files, cite_map)
    refs = add_references(doc, order, bib)
    set_core_properties(doc)
    doc.save(FINAL_DOCX)
    write_reports(eq_count, formulas, refs, order)
    print(json.dumps({"docx": str(FINAL_DOCX), "display_formula_count": eq_count,
                      "reference_count": len(refs), "citation_order": order},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
