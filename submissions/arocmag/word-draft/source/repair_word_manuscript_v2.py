from __future__ import annotations

import copy
import csv
import hashlib
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[4]
WD = ROOT / "submissions" / "arocmag" / "word-draft"
V1 = WD / "manuscript" / "kwaay-arocmag-draft.docx"
V2 = WD / "manuscript" / "kwaay-arocmag-draft-v2.docx"
TEMPLATE = WD / "source" / "official-template-converted.docx"
REF_JSON = WD / "references" / "citation-order.json"
REF_TSV = WD / "references" / "citation-order.tsv"
REPORTS = WD / "reports"
ACCESS_DATE = "2026-10-06"
HEADER_DATE = "2026/10/06"


REFERENCES = [
    {
        "number": 1,
        "key": "cohngordon2020signal",
        "text": "[1] Cohn-Gordon K, Cremers C, Dowling B, et al. A formal security analysis of the Signal messaging protocol[J]. Journal of Cryptology, 2020, 33(4): 1914-1983.",
        "needs_confirmation": False,
        "source": "Original BibTeX; Springer/Crossref DOI metadata (10.1007/s00145-020-09360-1).",
        "confirmed": "YES",
    },
    {
        "number": 2,
        "key": "cremers2023session",
        "text": "[2] Cremers C, Jacomme C, Naska A. Formal analysis of session-handling in secure messaging: lifting security from sessions to conversations[C]// Proc of the 32nd USENIX Security Symposium (USENIX Security 23). Berkeley, CA: USENIX Association, 2023: 1235-1252.",
        "needs_confirmation": False,
        "source": "Original BibTeX; USENIX official paper page/BibTeX; USENIX official Berkeley address.",
        "confirmed": "YES",
    },
    {
        "number": 3,
        "key": "bhargavan2024pqxdh",
        "text": "[3] Bhargavan K, Jacomme C, Kiefer F, et al. Formal verification of the PQXDH post-quantum key agreement protocol for end-to-end secure messaging[C]// Proc of the 33rd USENIX Security Symposium (USENIX Security 24). Berkeley, CA: USENIX Association, 2024: 469-486.",
        "needs_confirmation": False,
        "source": "Original BibTeX; USENIX official paper page/BibTeX; USENIX official Berkeley address.",
        "confirmed": "YES",
    },
    {
        "number": 4,
        "key": "collins2024kwaay",
        "text": "[4] Collins D, Huguenin-Dumittan L, Nguyen N K, et al. K-Waay: fast and deniable post-quantum X3DH without ring signatures[C]// Proc of the 33rd USENIX Security Symposium (USENIX Security 24). Berkeley, CA: USENIX Association, 2024: 433-450.",
        "needs_confirmation": False,
        "source": "Original BibTeX; USENIX official paper page/BibTeX; USENIX official Berkeley address.",
        "confirmed": "YES",
    },
    {
        "number": 5,
        "key": "collins2024kwaayfull",
        "text": "[5] Collins D, Huguenin-Dumittan L, Nguyen N K, et al. K-Waay: fast and deniable post-quantum X3DH without ring signatures[EB/OL]. Cryptology ePrint Archive, Paper 2024/120, 2024[2026-10-06]. https://eprint.iacr.org/2024/120.",
        "needs_confirmation": False,
        "source": "Original BibTeX; IACR Cryptology ePrint Archive official record and paper.",
        "confirmed": "YES",
    },
    {
        "number": 6,
        "key": "lupetti2006names",
        "text": "[6] Lupetti S, Dillema F W, Stabell-Kulø T. Names in cryptographic protocols[C]// Proc of the 4th International Workshop on Security in Information Systems (WOSIS). Setúbal, Portugal: SciTePress, 2006: 185-194.",
        "needs_confirmation": False,
        "source": "Original BibTeX; SciTePress official paper record (DOI 10.5220/0002484701850194); SciTePress official contact page.",
        "confirmed": "YES",
    },
    {
        "number": 7,
        "key": "lowe1997hierarchy",
        "text": "[7] Lowe G. A hierarchy of authentication specifications[C]// Proc of the 10th IEEE Computer Security Foundations Workshop. Los Alamitos, CA: IEEE Computer Society Press, 1997: 31-43.",
        "needs_confirmation": False,
        "source": "Original BibTeX; IEEE/Crossref DOI metadata (10.1109/CSFW.1997.596782); IEEE Computer Society official publications-office address.",
        "confirmed": "YES",
    },
    {
        "number": 8,
        "key": "meier2013tamarin",
        "text": "[8] Meier S, Schmidt B, Cremers C, et al. The TAMARIN prover for the symbolic analysis of security protocols[C]// Proc of the 25th International Conference on Computer Aided Verification (CAV), LNCS 8044. Berlin: Springer, 2013: 696-701.",
        "needs_confirmation": False,
        "source": "Original BibTeX; Springer/Crossref DOI metadata (10.1007/978-3-642-39799-8_48).",
        "confirmed": "YES",
    },
    {
        "number": 9,
        "key": "marlinspike2016x3dh",
        "text": "[9] Marlinspike M, Perrin T. The X3DH key agreement protocol[EB/OL]. Revision 1, 2016-11-04[2026-10-06]. https://signal.org/docs/specifications/x3dh/x3dh.pdf.",
        "needs_confirmation": False,
        "source": "Original BibTeX; Signal official X3DH specification PDF, revision line and title page.",
        "confirmed": "YES",
    },
    {
        "number": 10,
        "key": "wallez2023treesync",
        "text": "[10] Wallez T, Protzenko J, Beurdouche B, et al. TreeSync: authenticated group management for Messaging Layer Security[C]// Proc of the 32nd USENIX Security Symposium (USENIX Security 23). Berkeley, CA: USENIX Association, 2023: 1217-1233.",
        "needs_confirmation": False,
        "source": "Original BibTeX; USENIX official paper page/BibTeX; USENIX official Berkeley address.",
        "confirmed": "YES",
    },
    {
        "number": 11,
        "key": "balbas2023administration",
        "text": "[11] Balbás D, Collins D, Vaudenay S. Cryptographic administration for secure group messaging[C]// Proc of the 32nd USENIX Security Symposium (USENIX Security 23). Berkeley, CA: USENIX Association, 2023: 1253-1270.",
        "needs_confirmation": False,
        "source": "Original BibTeX; USENIX official paper page/BibTeX; USENIX official Berkeley address.",
        "confirmed": "YES",
    },
    {
        "number": 12,
        "key": "alwen2021modular",
        "text": "[12] Alwen J, Coretti S, Dodis Y, et al. Modular design of secure group messaging protocols and the security of MLS[C]// Proc of the ACM SIGSAC Conference on Computer and Communications Security. New York: ACM, 2021: 1463-1483.",
        "needs_confirmation": False,
        "source": "Original BibTeX; ACM/Crossref DOI metadata (10.1145/3460120.3484820); ACM official New York address.",
        "confirmed": "YES",
    },
    {
        "number": 13,
        "key": "cremers2021healing",
        "text": "[13] Cremers C, Hale B, Kohbrok K. The complexities of healing in secure group messaging: why cross-group effects matter[C]// Proc of the 30th USENIX Security Symposium (USENIX Security 21). Berkeley, CA: USENIX Association, 2021: 1847-1864.",
        "needs_confirmation": False,
        "source": "Original BibTeX; USENIX official paper page/BibTeX; USENIX official Berkeley address.",
        "confirmed": "YES",
    },
    {
        "number": 14,
        "key": "unger2015sok",
        "text": "[14] Unger N, Dechand S, Bonneau J, et al. SoK: secure messaging[C]// Proc of IEEE Symposium on Security and Privacy. [S. l. ]: IEEE, 2015: 232-249.",
        "needs_confirmation": True,
        "source": "Original BibTeX; IEEE Xplore official record; Crossref DOI metadata (10.1109/SP.2015.22). The publisher is confirmed, but the proceedings publication place is not directly stated.",
        "confirmed": "NEEDS_CONFIRMATION",
    },
    {
        "number": 15,
        "key": "andova2008compositional",
        "text": "[15] Andova S, Cremers C, Gjøsteen K, et al. A framework for compositional verification of security protocols[J]. Information and Computation, 2008, 206(2-4): 425-459.",
        "needs_confirmation": False,
        "source": "Original BibTeX; Elsevier/Crossref DOI metadata (10.1016/j.ic.2007.07.002).",
        "confirmed": "YES",
    },
    {
        "number": 16,
        "key": "ceelen2008chosenname",
        "text": "[16] Ceelen P, Mauw S, Radomirović S. Chosen-name attacks: an overlooked class of type-flaw attacks[J]. Electronic Notes in Theoretical Computer Science, 2008, 197(2): 31-43.",
        "needs_confirmation": False,
        "source": "Original BibTeX; Elsevier/Crossref DOI metadata (10.1016/j.entcs.2007.12.015).",
        "confirmed": "YES",
    },
    {
        "number": 17,
        "key": "peltonen2020misbinding",
        "text": "[17] Peltonen A, Sethi M, Aura T. Formal verification of misbinding attacks on secure device pairing and bootstrapping[J]. Journal of Information Security and Applications, 2020, 51: 102461.",
        "needs_confirmation": False,
        "source": "Original BibTeX; Elsevier/Crossref DOI metadata (10.1016/j.jisa.2020.102461).",
        "confirmed": "YES",
    },
    {
        "number": 18,
        "key": "thomson2021uks",
        "text": "[18] Thomson M, Rescorla E. RFC 8844, Unknown key-share attacks on uses of TLS with the Session Description Protocol (SDP)[S/OL]. [S. l. ]: RFC Editor, 2021[2026-10-06]. https://www.rfc-editor.org/rfc/rfc8844.html.",
        "needs_confirmation": True,
        "source": "Original BibTeX; RFC Editor official RFC 8844 record; DOI 10.17487/RFC8844. The standards-track status and publisher are confirmed, but a proceedings-style publication place is not stated.",
        "confirmed": "NEEDS_CONFIRMATION",
    },
    {
        "number": 19,
        "key": "sattarzadeh2015injective",
        "text": "[19] Sattarzadeh B, Fallah M S. Automated type-based analysis of injective agreement in the presence of compromised principals[J]. Journal of Logical and Algebraic Methods in Programming, 2015, 84(5): 576-610.",
        "needs_confirmation": False,
        "source": "Original BibTeX; Elsevier/Crossref DOI metadata (10.1016/j.jlamp.2015.06.002).",
        "confirmed": "YES",
    },
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def replace_story(story, source_story) -> None:
    dst = story._element
    for child in list(dst):
        dst.remove(child)
    for child in source_story._element:
        dst.append(copy.deepcopy(child))


def replace_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def insert_after(paragraph, source_paragraph, text: str):
    node = copy.deepcopy(source_paragraph._p)
    paragraph._p.addnext(node)
    from docx.text.paragraph import Paragraph

    inserted = Paragraph(node, paragraph._parent)
    replace_text(inserted, text)
    return inserted


def repair_docx() -> None:
    shutil.copy2(V1, V2)
    doc = Document(V2)
    template = Document(TEMPLATE)

    # Front matter: remove the body-level funding placeholder, then insert the two
    # omitted official-template fields without changing title, abstract or authors.
    fund_body = next(
        p for p in doc.paragraphs
        if p.text.strip() == "基金项目、作者简介及通信作者信息待补充并核验"
    )
    fund_body._element.getparent().remove(fund_body._element)

    zh_keywords = next(p for p in doc.paragraphs if p.text.startswith("关键词："))
    insert_after(zh_keywords, template.paragraphs[5], "中图分类号：待确认")
    en_authors = next(p for p in doc.paragraphs if p.text.startswith("Author A,"))
    insert_after(
        en_authors,
        template.paragraphs[8],
        "(Department / Institution / City / Postal Code / China — to be completed and verified)",
    )

    # Restore the official default/first-page header and first-page footer stories.
    first = doc.sections[0]
    src = template.sections[0]
    first.different_first_page_header_footer = True
    first.header.is_linked_to_previous = False
    first.first_page_header.is_linked_to_previous = False
    first.footer.is_linked_to_previous = False
    first.first_page_footer.is_linked_to_previous = False
    replace_story(first.header, src.header)
    replace_story(first.first_page_header, src.first_page_header)
    replace_story(first.footer, src.footer)
    replace_story(first.first_page_footer, src.first_page_footer)

    for p in first.first_page_header.paragraphs:
        if "修改日期" in p.text:
            replace_text(p, p.text.replace("2021/07/27", HEADER_DATE))

    footer = first.first_page_footer
    while len(footer.paragraphs) < 4:
        footer._element.append(copy.deepcopy(footer.paragraphs[-1]._p))
    replace_text(footer.paragraphs[0], "——————————")
    replace_text(footer.paragraphs[1], "　　基金项目：待作者确认")
    replace_text(footer.paragraphs[2], "　　作者简介：待作者确认")
    replace_text(footer.paragraphs[3], "　　通信作者及电子邮箱：待作者确认")

    for section in doc.sections[1:]:
        section.different_first_page_header_footer = False
        section.header.is_linked_to_previous = True
        section.first_page_header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True
        section.first_page_footer.is_linked_to_previous = True

    # Replace the 19 reference paragraphs in place so numbering, style and section
    # geometry remain unchanged.
    ref_heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "参考文献")
    ref_paragraphs = doc.paragraphs[ref_heading_index + 1: ref_heading_index + 20]
    if len(ref_paragraphs) != 19:
        raise RuntimeError(f"Expected 19 reference paragraphs, found {len(ref_paragraphs)}")
    for paragraph, item in zip(ref_paragraphs, REFERENCES, strict=True):
        replace_text(paragraph, re.sub(r"^\[\d+\]\s*", "", item["text"]))

    # Preserve anonymity and remove accidental identifying core properties.
    doc.core_properties.author = "Anonymous Author(s)"
    doc.core_properties.last_modified_by = "Anonymous Author(s)"
    doc.save(V2)


def write_reference_maps() -> None:
    REF_JSON.write_text(json.dumps([
        {k: item[k] for k in ("number", "key", "text", "needs_confirmation")}
        for item in REFERENCES
    ], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with REF_TSV.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, delimiter="\t", lineterminator="\n")
        writer.writerow(["number", "key", "needs_confirmation", "reference"])
        for item in REFERENCES:
            writer.writerow([
                item["number"], item["key"], str(item["needs_confirmation"]).lower(), item["text"]
            ])


def write_reference_audit() -> None:
    lines = [
        "# REFERENCE AUDIT V2",
        "",
        "核验口径：作者姓名、题名、文献类型、载体、年份、卷期、页码、出版者、DOI/URL 均以原始 BibTeX 与出版社、DOI 注册机构或规范发布者的一手记录交叉核对。会议条目的出版地按出版者所在地著录，不采用会议举办地。外文作者按期刊规则写为姓在前、名首字母大写且不加点；超过三人列前三人后加 `et al`。题名采用句首大写，专名与缩略语保留原有大小写。",
        "",
        "| 编号 | BibTeX key | 最终格式 | 信息来源 | 完全确认 |",
        "|---:|---|---|---|---|",
    ]
    for item in REFERENCES:
        text = item["text"].replace("|", "\\|")
        source = item["source"].replace("|", "\\|")
        lines.append(f"| {item['number']} | `{item['key']}` | {text} | {source} | {item['confirmed']} |")
    lines += [
        "",
        "## 仍需确认",
        "",
        "- `[14] unger2015sok`：IEEE 官方记录确认作者、题名、会议、年份、页码、DOI 与出版者；未发现直接给出该会议录出版地的一手出版项，故依期刊规则著录 `[S. l. ]`。",
        "- `[18] thomson2021uks`：RFC Editor 官方记录确认标准号、题名、作者、发布日期、DOI 与 URL；未发现适用于该在线标准的明确出版地，故著录 `[S. l. ]`。",
        "",
        "除上述出版地字段外，其余 17 条均已由一手记录或 DOI 元数据确认。`needs_confirmation` 已同步写入 JSON/TSV 映射。",
    ]
    (REPORTS / "REFERENCE_AUDIT_V2.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_formula_audit() -> None:
    lines = [
        "# FORMULA AUDIT V2",
        "",
        "## 状态",
        "",
        "MATH_TYPE_CONVERSION_BLOCKED",
        "",
        "当前 Windows 环境未发现可调用的 MathType：常见 Program Files、ProgramData、用户 AppData 路径无 MathType/Design Science/WIRIS 文件，卸载注册表、Word Add-ins 注册表、命令入口与相关进程也均无匹配项。因此未伪造 MathType OLE 对象，也未把公式转成图片。",
        "",
        "## 保留项",
        "",
        "- v2 保留 197 个可编辑 OMML 数学对象。",
        "- 27 个编号行间公式及其编号保持不变。",
        "- `source/formulas.json` 保留且未改写。",
        "- 量词与关系符号（∀、∃、⇒、⇔、≠）、上下标、时间点和多行公式继续由 OMML 表示。",
        "",
        "由于期刊要求 MathType 而当前环境无法执行真实转换，本稿不能标记为投稿就绪。",
    ]
    (REPORTS / "FORMULA_AUDIT_V2.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    REPORTS.mkdir(parents=True, exist_ok=True)
    repair_docx()
    write_reference_maps()
    write_reference_audit()
    write_formula_audit()
    cache = WD / "source" / "__pycache__"
    if cache.exists():
        shutil.rmtree(cache)
    print(json.dumps({
        "v2": str(V2),
        "sha256": sha256(V2),
        "references": len(REFERENCES),
        "needs_confirmation": [x["number"] for x in REFERENCES if x["needs_confirmation"]],
        "pycache_removed": not cache.exists(),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
